import pandas as pd
import numpy as np
import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from datetime import datetime
import sys
import warnings

warnings.filterwarnings('ignore')


class TimesheetProcessor:
    def __init__(self):
        self.timesheet_df = None
        self.filter_conditions = {}

    def load_timesheet(self, file_path):
        """加载工时记录表"""
        try:
            if not os.path.exists(file_path):
                raise FileNotFoundError(f"文件不存在: {file_path}")

            self.timesheet_df = pd.read_excel(file_path, engine='openpyxl')
            self.timesheet_df = self.optimize_data_types(self.timesheet_df)
            return True

        except Exception as e:
            raise Exception(f"加载文件失败: {str(e)}")

    def optimize_data_types(self, df):
        """优化数据类型减少内存占用"""
        df_opt = df.copy()

        if 'Hours' in df_opt.columns:
            df_opt['Hours'] = pd.to_numeric(df_opt['Hours'], errors='coerce', downcast='float')

        text_cols = ['Client Name', 'Engagement Code', 'Employee Name', 'COE Manager', 'Description']
        for col in text_cols:
            if col in df_opt.columns and not df_opt[col].empty:
                unique_ratio = df_opt[col].nunique() / len(df_opt) if len(df_opt) > 0 else 1
                if unique_ratio < 0.5:
                    df_opt[col] = df_opt[col].astype('category')

        return df_opt

    def load_filter_conditions(self, filter_files, filter_columns):
        """加载筛选条件"""
        self.filter_conditions = {}

        for i, (file_path, filter_column) in enumerate(zip(filter_files, filter_columns)):
            if not file_path or not filter_column:
                continue

            try:
                filter_df = pd.read_excel(file_path, engine='openpyxl')
                if len(filter_df) > 0:
                    filter_values = set(filter_df.iloc[:, 0].dropna().astype(str).str.strip())
                    self.filter_conditions[filter_column] = filter_values
            except Exception as e:
                print(f"加载条件文件失败: {str(e)}")

    def extract_order_number(self, description):
        """从Description列提取订单号"""
        if pd.isna(description):
            return None

        description_str = str(description).strip()

        patterns = [
            r'^(\d{8})(?=\D|$)',
            r'(\d{8})[\s\-_]',
            r'(\d{8})[^\d]',
            r'\b(\d{8})\b',
        ]

        for pattern in patterns:
            match = re.search(pattern, description_str)
            if match:
                order_num = match.group(1)
                if len(order_num) == 8 and order_num.isdigit():
                    return order_num

        return None

    def apply_filters(self, logic='and'):
        """应用筛选条件"""
        if self.timesheet_df is None or len(self.timesheet_df) == 0:
            return pd.DataFrame()

        if not self.filter_conditions:
            return self.timesheet_df.copy()

        conditions = []

        for filter_column, filter_values in self.filter_conditions.items():
            if filter_column not in self.timesheet_df.columns:
                continue

            if filter_column == 'Description':
                self.timesheet_df['Extracted_Order'] = self.timesheet_df[filter_column].apply(
                    self.extract_order_number
                )
                condition = self.timesheet_df['Extracted_Order'].isin(filter_values)
            else:
                condition = self.timesheet_df[filter_column].astype(str).str.strip().isin(filter_values)

            conditions.append(condition)

        if conditions:
            if logic.lower() == 'and':
                final_condition = conditions[0]
                for condition in conditions[1:]:
                    final_condition = final_condition & condition
            else:
                final_condition = conditions[0]
                for condition in conditions[1:]:
                    final_condition = final_condition | condition

            filtered_df = self.timesheet_df[final_condition].copy()

            if 'Extracted_Order' in filtered_df.columns:
                filtered_df = filtered_df.drop('Extracted_Order', axis=1)

            return filtered_df
        else:
            return self.timesheet_df.copy()

    def safe_fillna_categorical(self, series, fill_value='Unknown'):
        """安全地填充分类列的空值"""
        if pd.api.types.is_categorical_dtype(series):
            if fill_value not in series.cat.categories:
                series = series.cat.add_categories([fill_value])
            return series.fillna(fill_value)
        else:
            return series.fillna(fill_value)

    def generate_reports(self, filtered_df, has_description_filter=False):
        """生成报表"""
        try:
            if filtered_df is None or len(filtered_df) == 0:
                empty_summary = pd.DataFrame(columns=['Client Name', 'Engagement Code', 'Hours'])
                empty_detail = pd.DataFrame(
                    columns=['Client Name', 'Engagement Code', 'Employee Name', 'Hours', 'Transaction Cycle Date'])
                empty_order = pd.DataFrame(columns=['订单号', '原始描述', '工时合计'])
                return empty_summary, empty_detail, empty_order, pd.DataFrame()

            filter_result_df = filtered_df.copy()

            groupby_columns = []
            for col in ['Client Name', 'Engagement Code']:
                if col in filtered_df.columns:
                    groupby_columns.append(col)

            if groupby_columns:
                for col in groupby_columns:
                    if filtered_df[col].isna().any():
                        filtered_df[col] = self.safe_fillna_categorical(filtered_df[col], 'Unknown')

                summary_df = filtered_df.groupby(groupby_columns, as_index=False, observed=True)['Hours'].sum()
            else:
                summary_total = filtered_df['Hours'].sum()
                summary_df = pd.DataFrame({'Hours': [summary_total]})

            detail_columns = ['Client Name', 'Engagement Code', 'Employee Name', 'Hours', 'Transaction Cycle Date']
            available_columns = [col for col in detail_columns if col in filtered_df.columns]
            detail_df = filtered_df[available_columns].copy()

            if has_description_filter and 'Description' in filtered_df.columns:
                filtered_df['订单号'] = filtered_df['Description'].apply(self.extract_order_number)
                order_summary_df = filtered_df[filtered_df['订单号'].notna()].groupby(
                    '订单号', as_index=False, observed=True
                )['Hours'].sum()

                order_descriptions = filtered_df.groupby('订单号', observed=True)['Description'].first().reset_index()
                order_summary_df = order_summary_df.merge(order_descriptions, on='订单号', how='left')
                order_summary_df = order_summary_df[['订单号', 'Description', 'Hours']]
                order_summary_df.columns = ['订单号', '原始描述', '工时合计']
                order_summary_df = order_summary_df.sort_values('工时合计', ascending=False)
            else:
                order_summary_df = pd.DataFrame(columns=['订单号', '原始描述', '工时合计'])

            if not self.validate_data_consistency_strict(filter_result_df, summary_df, detail_df):
                raise Exception("数据一致性验证失败")

            return summary_df, detail_df, order_summary_df, filter_result_df

        except Exception as e:
            raise Exception(f"生成报表失败: {str(e)}")

    def validate_data_consistency_strict(self, filter_result_df, summary_df, detail_df):
        """严格的数据一致性验证"""
        try:
            if len(filter_result_df) == 0:
                return True

            filter_total = filter_result_df['Hours'].sum()
            summary_total = summary_df['Hours'].sum()
            detail_total = detail_df['Hours'].sum()

            tolerance = 1e-10

            issues = []
            if abs(filter_total - summary_total) > tolerance:
                issues.append(f"筛选结果与汇总表不一致: {filter_total:.6f} vs {summary_total:.6f}")

            if abs(filter_total - detail_total) > tolerance:
                issues.append(f"筛选结果与明细表不一致: {filter_total:.6f} vs {detail_total:.6f}")

            if issues:
                for issue in issues:
                    print(f"❌ {issue}")
                return False
            else:
                print("✅ 数据一致性验证通过")
                return True

        except Exception as e:
            print(f"数据一致性验证错误: {e}")
            return False


class CompactTimesheetGUI:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("工时统计系统 - 智能列选择版")
        self.root.geometry("800x650+100+50")
        self.root.minsize(750, 600)

        self.timesheet_path = tk.StringVar()
        self.filter_files = [tk.StringVar() for _ in range(3)]
        self.filter_columns = [tk.StringVar(value="") for _ in range(3)]
        self.logic_var = tk.StringVar(value="and")
        self.current_condition = tk.IntVar(value=0)

        self.processor = TimesheetProcessor()
        self.available_columns = []
        self.filter_file_columns = [[] for _ in range(3)]  # 存储每个筛选文件的列名

        self.setup_ui()

    def setup_ui(self):
        """设置用户界面"""
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)

        title_frame = ttk.Frame(main_frame)
        title_frame.pack(fill=tk.X, pady=(0, 10))

        title_label = ttk.Label(title_frame, text="工时统计系统", font=("Arial", 16, "bold"))
        title_label.pack(pady=5)

        desc_label = ttk.Label(title_frame, text="智能列选择 • 数据一致性 • 紧凑界面", font=("Arial", 10))
        desc_label.pack(pady=2)

        self.setup_action_buttons(main_frame)

        notebook = ttk.Notebook(main_frame)
        notebook.pack(fill=tk.BOTH, expand=True, pady=10)

        file_tab = self.create_file_tab(notebook)
        notebook.add(file_tab, text="文件选择")

        filter_tab = self.create_filter_tab(notebook)
        notebook.add(filter_tab, text="筛选条件")

        self.setup_status_log_area(main_frame)

    def setup_action_buttons(self, parent):
        """设置操作按钮区域"""
        action_frame = ttk.Frame(parent)
        action_frame.pack(fill=tk.X, pady=(0, 10))

        button_container = ttk.Frame(action_frame)
        button_container.pack(fill=tk.X, pady=5)

        button_container.columnconfigure(0, weight=1)
        button_container.columnconfigure(1, weight=1)
        button_container.columnconfigure(2, weight=1)

        self.generate_button = ttk.Button(button_container, text="生成报表",
                                          command=self.generate_report, width=12)
        self.generate_button.grid(row=0, column=0, padx=5, pady=2, sticky="ew")

        ttk.Button(button_container, text="清空选择",
                   command=self.clear_selection, width=10).grid(row=0, column=1, padx=5, pady=2, sticky="ew")

        ttk.Button(button_container, text="退出",
                   command=self.root.quit, width=8).grid(row=0, column=2, padx=5, pady=2, sticky="ew")

    def create_file_tab(self, parent):
        """创建文件选择标签页"""
        tab = ttk.Frame(parent, padding=10)

        file_frame = ttk.LabelFrame(tab, text="基础工时记录表", padding=15)
        file_frame.pack(fill=tk.X, pady=10)

        ttk.Label(file_frame, text="文件路径:").pack(anchor=tk.W)

        entry_frame = ttk.Frame(file_frame)
        entry_frame.pack(fill=tk.X, pady=10)

        entry = ttk.Entry(entry_frame, textvariable=self.timesheet_path)
        entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        ttk.Button(entry_frame, text="选择文件",
                   command=self.select_timesheet).pack(side=tk.RIGHT, padx=5)

        logic_frame = ttk.LabelFrame(tab, text="筛选逻辑", padding=15)
        logic_frame.pack(fill=tk.X, pady=10)

        ttk.Radiobutton(logic_frame, text="AND（同时满足所有条件）",
                        variable=self.logic_var, value="and").pack(anchor=tk.W)
        ttk.Radiobutton(logic_frame, text="OR（满足任一条件即可）",
                        variable=self.logic_var, value="or").pack(anchor=tk.W)

        return tab

    def create_filter_tab(self, parent):
        """创建筛选条件标签页"""
        tab = ttk.Frame(parent, padding=10)

        slider_frame = ttk.LabelFrame(tab, text="条件选择", padding=10)
        slider_frame.pack(fill=tk.X, pady=10)

        ttk.Label(slider_frame, text="滑动选择条件 (1-3):").pack(anchor=tk.W)

        slider_container = ttk.Frame(slider_frame)
        slider_container.pack(fill=tk.X, pady=5)

        ttk.Label(slider_container, text="1", font=("Arial", 8)).pack(side=tk.LEFT)

        condition_slider = tk.Scale(
            slider_container,
            from_=0, to=2,
            variable=self.current_condition,
            orient=tk.HORIZONTAL,
            length=300,
            showvalue=True,
            resolution=1,
            command=self.on_slider_change
        )
        condition_slider.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)

        ttk.Label(slider_container, text="3", font=("Arial", 8)).pack(side=tk.RIGHT)

        self.condition_label = ttk.Label(slider_frame, text="当前配置: 条件 1",
                                         font=("Arial", 10))
        self.condition_label.pack(pady=5)

        self.condition_config_frame = ttk.LabelFrame(tab, text="条件配置", padding=10)
        self.condition_config_frame.pack(fill=tk.X, pady=10)
        self.update_condition_display()

        return tab

    def setup_status_log_area(self, parent):
        """设置状态和日志区域"""
        self.status_var = tk.StringVar(value="就绪")
        status_label = ttk.Label(parent, textvariable=self.status_var,
                                 relief=tk.SUNKEN, padding=8)
        status_label.pack(fill=tk.X, pady=5)

        self.progress = ttk.Progressbar(parent, mode='indeterminate')
        self.progress.pack(fill=tk.X, pady=5)

        log_frame = ttk.LabelFrame(parent, text="处理日志", padding=5)
        log_frame.pack(fill=tk.BOTH, expand=True, pady=10)

        text_frame = ttk.Frame(log_frame)
        text_frame.pack(fill=tk.BOTH, expand=True)

        self.log_text = tk.Text(text_frame, height=12, font=("Consolas", 8))
        scrollbar = ttk.Scrollbar(text_frame, orient="vertical", command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scrollbar.set)

        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def on_slider_change(self, value):
        """滑动条值改变时的回调"""
        condition_index = int(float(value))
        self.current_condition.set(condition_index)
        self.condition_label.config(text=f"当前配置: 条件 {condition_index + 1}")
        self.update_condition_display()

    def update_condition_display(self):
        """更新条件配置显示"""
        for widget in self.condition_config_frame.winfo_children():
            widget.destroy()

        condition_index = self.current_condition.get()

        file_frame = ttk.Frame(self.condition_config_frame)
        file_frame.pack(fill=tk.X, pady=5)

        ttk.Label(file_frame, text="条件文件:", width=8).pack(side=tk.LEFT)
        entry = ttk.Entry(file_frame, textvariable=self.filter_files[condition_index])
        entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        ttk.Button(file_frame, text="选择",
                   command=lambda idx=condition_index: self.select_filter_file(idx),
                   width=6).pack(side=tk.RIGHT, padx=5)

        column_frame = ttk.Frame(self.condition_config_frame)
        column_frame.pack(fill=tk.X, pady=5)

        ttk.Label(column_frame, text="匹配列:", width=8).pack(side=tk.LEFT)

        # 使用Combobox替代Entry，自动从筛选文件读取列名
        self.column_combo = ttk.Combobox(
            column_frame,
            textvariable=self.filter_columns[condition_index],
            values=self.filter_file_columns[condition_index],
            width=25,
            state="readonly"  # 设置为只读，只能从下拉列表选择
        )
        self.column_combo.pack(side=tk.LEFT, padx=5)

        # 添加刷新按钮，用于重新读取列名
        ttk.Button(column_frame, text="刷新列名",
                   command=lambda idx=condition_index: self.refresh_filter_columns(idx),
                   width=8).pack(side=tk.RIGHT, padx=5)

        help_text = """💡 说明：选择条件文件后，点击"刷新列名"按钮加载可用列，然后从下拉列表选择匹配列"""
        help_label = ttk.Label(self.condition_config_frame, text=help_text,
                               justify=tk.LEFT, font=("Arial", 8))
        help_label.pack(anchor=tk.W, pady=8)

    def refresh_filter_columns(self, index):
        """刷新筛选文件的列名列表"""
        file_path = self.filter_files[index].get()
        if file_path and os.path.exists(file_path):
            try:
                sample_df = pd.read_excel(file_path, nrows=5, engine='openpyxl')
                self.filter_file_columns[index] = list(sample_df.columns)
                self.column_combo['values'] = self.filter_file_columns[index]
                self.log_message(f"条件{index + 1}文件列名已刷新: {len(self.filter_file_columns[index])} 列")
            except Exception as e:
                self.log_message(f"刷新条件{index + 1}文件列名失败: {str(e)}")
        else:
            self.log_message(f"条件{index + 1}文件不存在或未选择")

    def update_columns_list(self):
        """更新基础文件的可用列列表"""
        if self.timesheet_path.get() and os.path.exists(self.timesheet_path.get()):
            try:
                sample_df = pd.read_excel(self.timesheet_path.get(), nrows=5, engine='openpyxl')
                self.available_columns = list(sample_df.columns)
            except Exception as e:
                print(f"获取列名失败: {e}")

    def select_timesheet(self):
        """选择工时记录表"""
        filename = filedialog.askopenfilename(
            title="选择基础工时记录表",
            filetypes=[("Excel files", "*.xlsx;*.xls"), ("All files", "*.*")]
        )
        if filename:
            self.timesheet_path.set(filename)
            self.status_var.set(f"已选择: {os.path.basename(filename)}")
            self.log_message(f"选择基础文件: {filename}")
            self.update_columns_list()

    def select_filter_file(self, index):
        """选择筛选条件文件，并自动加载列名"""
        filename = filedialog.askopenfilename(
            title=f"选择条件{index + 1}文件",
            filetypes=[("Excel files", "*.xlsx;*.xls"), ("All files", "*.*")]
        )
        if filename:
            self.filter_files[index].set(filename)
            self.log_message(f"选择条件{index + 1}文件: {filename}")

            # 自动加载列名
            self.refresh_filter_columns(index)

    def log_message(self, message):
        """添加日志消息"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        log_entry = f"[{timestamp}] {message}\n"
        self.log_text.insert(tk.END, log_entry)
        self.log_text.see(tk.END)
        self.root.update()

    def clear_selection(self):
        """清空所有选择"""
        self.timesheet_path.set("")
        for i in range(3):
            self.filter_files[i].set("")
            self.filter_columns[i].set("")
            self.filter_file_columns[i] = []
        self.status_var.set("已清空")
        self.log_text.delete(1.0, tk.END)
        self.log_message("已清空所有选择")

    def generate_report(self):
        """生成报表"""
        try:
            if not self.timesheet_path.get():
                messagebox.showerror("错误", "请选择基础工时记录表")
                return

            self.generate_button.config(state='disabled')
            self.status_var.set("处理中...")
            self.progress.start(10)
            self.log_message("开始处理工时数据...")
            self.root.update()

            if not self.processor.load_timesheet(self.timesheet_path.get()):
                raise Exception("加载基础数据失败")

            filter_files = []
            filter_columns = []

            for i in range(3):
                if self.filter_files[i].get() and self.filter_columns[i].get():
                    filter_files.append(self.filter_files[i].get())
                    filter_columns.append(self.filter_columns[i].get())

            self.log_message(f"应用筛选条件: {len(filter_files)} 个")

            self.processor.load_filter_conditions(filter_files, filter_columns)

            filtered_df = self.processor.apply_filters(self.logic_var.get())

            if len(filtered_df) == 0:
                self.log_message("警告: 筛选后无数据")
                if not messagebox.askyesno("警告", "筛选后无数据，是否生成空报表？"):
                    return

            has_description_filter = any('Description' in col for col in self.processor.filter_conditions.keys())
            summary_df, detail_df, order_summary_df, filter_result_df = self.processor.generate_reports(
                filtered_df, has_description_filter
            )

            if summary_df is not None:
                output_file = self.save_results(summary_df, detail_df, order_summary_df, filter_result_df)
                self.status_var.set("完成!")
                self.log_message(f"报表保存至: {output_file}")
                messagebox.showinfo("成功", f"工时报表生成成功!\n文件位置: {output_file}")
            else:
                raise Exception("生成报表失败")

        except Exception as e:
            error_msg = f"处理失败: {str(e)}"
            self.status_var.set(f"错误")
            self.log_message(f"错误: {error_msg}")
            messagebox.showerror("错误", error_msg)
        finally:
            self.generate_button.config(state='normal')
            self.progress.stop()

    def save_results(self, summary_df, detail_df, order_summary_df, filter_result_df):
        """保存结果到Excel"""
        output_dir = os.path.dirname(self.timesheet_path.get())
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_file = os.path.join(output_dir, f"工时统计报表_{timestamp}.xlsx")

        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            if not filter_result_df.empty:
                filter_result_df.to_excel(writer, sheet_name='筛选结果', index=False)
            else:
                pd.DataFrame({'提示': ['无筛选结果数据']}).to_excel(writer, sheet_name='筛选结果', index=False)

            summary_df.to_excel(writer, sheet_name='汇总表', index=False)
            detail_df.to_excel(writer, sheet_name='明细表', index=False)

            if not order_summary_df.empty and len(order_summary_df) > 0:
                order_summary_df.to_excel(writer, sheet_name='订单汇总', index=False)
            else:
                pd.DataFrame({'提示': ['无订单汇总数据']}).to_excel(writer, sheet_name='订单汇总', index=False)

            workbook = writer.book
            for sheet_name in writer.sheets:
                worksheet = writer.sheets[sheet_name]
                worksheet.column_dimensions['A'].width = 15
                worksheet.column_dimensions['B'].width = 15
                worksheet.column_dimensions['C'].width = 15

        return output_file

    def run(self):
        """运行GUI"""
        self.root.mainloop()


def main():
    """主函数"""
    print("=== 工时统计系统 (智能列选择版) ===")
    print("优化内容:")
    print("1. ✅ 智能列选择 - 自动读取筛选文件列名")
    print("2. ✅ 下拉框选择 - 从列表选择匹配列")
    print("3. ✅ 自动刷新 - 选择文件后自动加载列名")
    print("4. ✅ 保持数据一致性 - 严格验证机制")

    try:
        import pandas as pd
        import openpyxl
        print(f"✅ Pandas版本: {pd.__version__}")
        print("✅ 依赖检查通过")
    except ImportError as e:
        print(f"❌ 缺少依赖: {e}")
        print("请运行: pip install pandas openpyxl")
        return

    app = CompactTimesheetGUI()
    app.run()


if __name__ == "__main__":
    main()
