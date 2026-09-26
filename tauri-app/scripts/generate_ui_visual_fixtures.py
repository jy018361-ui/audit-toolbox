"""Generate deterministic, fictional workbooks for UI screenshot walkthroughs.

The fixtures deliberately stress dense tables, long account names, multiple
entities, wide JE headers, missing values, and differing currency formats.
They contain no customer data and must not be used as accounting test oracles.
"""

from __future__ import annotations

from datetime import date, timedelta
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "fixtures" / "ui-visual"
PREVIEW_JSON = ROOT / "src" / "preview" / "demo" / "loanVisualFixture.json"

TB_HEADERS = [
    "GL Account Number", "GL Account Name", "Functional Beginning Balance",
    "Functional Ending Balance", "Functional Ending Balance Before Profit",
    "Functional Debit Activity Amount", "Functional Credit Activity Amount",
    "Functional Currency Code", "Business Unit", "Business Unit Name",
    "Account Level", "Last Level", "Fiscal Range", "Account Class",
    "Account sub class", "Account Type", "Account sub Type",
    "Document Beginning Balance", "Document Debit Activity Amount",
    "Document Credit Activity Amount", "Document Ending Balance",
    "Document Ending Balance Before Profit", "Document Currency Code", "Year",
]

JE_HEADERS = [
    "JE Number", "GL Account Number", "GL Account Name", "Debit_Credit",
    "Functional Debit Amount", "Functional Credit Amount", "Functional Amount",
    "Effective Date", "Entry Date", "Entry Time", "Preparer ID", "Preparer",
    "Source", "Source Name", "JE Description", "JE Line Description",
    "Year", "Period", "Functional Currency Code", "Business Unit",
    "Business Unit Name", "JE Line Number", "Document Debit Amount",
    "Document Credit Amount", "Document Amount",
] + [f"Additional Audit Field {index:02d}" for index in range(26, 78)]

ENTITIES = [
    ("SLP", "星澜项目管理有限公司"),
    ("HQ", "星澜集团总部"),
    ("EAST", "星澜华东事业部"),
]

ACCOUNTS = [
    ("200101", "短期借款-工商银行园区支行", "loan"),
    ("250101", "长期借款-股东借款(美元)", "loan"),
    ("660301", "财务费用-利息支出", "interest"),
    ("100101", "银行存款-中国银行苏州分行(美元户)", "other"),
    ("112301", "应收账款-大型工程项目跨区域结算及年度清算往来款", "other"),
    ("160401", "固定资产-数据中心服务器及配套网络设备", "other"),
]


def new_sheet(title: str, headers: list[str]):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}1"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEEEB")
    return book, sheet


def build_tb() -> Path:
    book, sheet = new_sheet("TB", TB_HEADERS)
    preview_accounts = []
    for index in range(120):
        entity_code, entity_name = ENTITIES[index % len(ENTITIES)]
        base_code, base_name, role = ACCOUNTS[index % len(ACCOUNTS)]
        suffix = index // len(ACCOUNTS) + 1
        code = f"{base_code}{suffix:03d}"
        name = f"{base_name}-第{suffix}项目"
        if index % 11 == 0:
            name += "-二〇二六年度专项审计调整及跨期重分类明细"
        currency = "USD" if index % 17 == 0 else "CNY"
        beginning = (2_900_000 + index * 53_017) * (-1 if role == "other" and index % 4 == 0 else 1)
        debit = 30_000 + index * 1_127
        credit = 55_000 + index * 1_033
        ending = beginning + debit - credit
        row = [
            code, name, beginning, ending, ending, debit, credit, currency,
            entity_code, entity_name, 3, "Y", "2026-01-01~2026-12-31",
            "负债" if role == "loan" else "损益" if role == "interest" else "资产",
            "审计样例", "明细科目", "测试", beginning, debit, credit, ending,
            ending, currency, 2026,
        ]
        if index % 19 == 0:
            row[5] = None  # Missing value exercises empty-cell layout.
        sheet.append(row)
        preview_accounts.append({
            "key": code,
            "identity": code,
            "code": code,
            "name": name,
            "currency": currency,
            "account": f"{code}-{name}",
            "opening": beginning,
            "closing": ending,
            "occurrence": debit + credit,
            "byEntity": [{"entity": entity_code, "opening": beginning, "closing": ending}],
            "suggestedType": "loan" if role == "loan" else "interest_expense" if role == "interest" else "skip",
        })
    path = OUT / "虚构_TB_120科目_长名称多主体.xlsx"
    book.save(path)
    PREVIEW_JSON.write_text(json.dumps(preview_accounts, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def build_je() -> Path:
    book, sheet = new_sheet("JE", JE_HEADERS)
    for index in range(900):
        entity_code, entity_name = ENTITIES[index % len(ENTITIES)]
        base_code, base_name, _ = ACCOUNTS[index % len(ACCOUNTS)]
        suffix = index // len(ACCOUNTS) % 20 + 1
        code = f"{base_code}{suffix:03d}"
        amount = 1200 + index * 37.15
        debit = amount if index % 2 == 0 else 0
        credit = amount if index % 2 else 0
        when = date(2026, 1, 1) + timedelta(days=index % 365)
        description = f"第{index + 1}笔虚构凭证：{base_name}的资金划拨与利息结转"
        if index % 13 == 0:
            description += "，涉及跨期、跨主体、外币折算及补充协议复核事项"
        row = [
            f"VISUAL-{index // 2 + 1:05d}", code, base_name,
            "D" if debit else "C", debit, credit, debit - credit,
            when, when, "09:30:00", f"U{index % 21:03d}", "虚构制单人",
            "视觉验收", "UI 测试凭证", description, description,
            2026, when.month, "USD" if index % 17 == 0 else "CNY",
            entity_code, entity_name, index % 4 + 1, debit, credit,
            debit - credit,
        ]
        row.extend(f"样例 {index + 1}-{column}" if column % 9 == 0 else None for column in range(26, 78))
        sheet.append(row)
    path = OUT / "虚构_JE_900行_77列_长摘要.xlsx"
    book.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for path in (build_tb(), build_je()):
        print(f"{path} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
