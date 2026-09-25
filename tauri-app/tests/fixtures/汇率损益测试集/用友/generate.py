# -*- coding: utf-8 -*-
"""
用友(U8/T+/T3)导出风格 汇率损益测试数据集 生成 + 自检 一体脚本

虚构主体: 友联进出口贸易有限公司 (本位币 CNY, 期间 2026-01-01 至 2026-06-30)
输出:
    用友_科目余额表.xlsx  -- 用友 U8 科目余额表风格 (标题行+单位信息行+列组表头)
    用友_序时账.xlsx      -- 用友 明细账/序时账 风格 (含汇率列/借贷原币本位币/余额)

预埋审计测试点 (详见 README.md):
    1. 期末未重估: 外币货币性科目(1002.02/1122.02 等)期末余额全部按历史记账汇率滚出,
       6月底无按期末汇率重估凭证, 6603.02 无期末重估发生额(仅已实现结汇/购汇价差)。
    2. 币种大小写混用: 序时账币种列混入 "usd"/"Usd" 共 3 行, 测试币种归一化。

控制台输出不使用 emoji (Windows GBK 兼容), 自检结果逐项打印, 任一失败退出码非 0。
复现: python generate.py
"""
import calendar
import os
import sys
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
ENTITY = "友联进出口贸易有限公司"

# 记账汇率(虚构中间价) 与 2026-06-30 期末汇率
BOOK = {"CNY": 1.00, "USD": 7.12, "EUR": 7.83, "HKD": 0.91, "JPY": 0.0482}
END_RATE = {"CNY": 1.00, "USD": 7.18, "EUR": 7.90, "HKD": 0.915, "JPY": 0.0495}
CNAME = {"CNY": "人民币", "USD": "美元", "EUR": "欧元", "HKD": "港币", "JPY": "日元"}

# 科目编码 -> (科目名称, 上级科目, 借贷性质 +1借/-1贷, 核算币种)
LEAVES = {
    "1002.01": ("人民币户", "1002", 1, "CNY"),
    "1002.02": ("美元户", "1002", 1, "USD"),
    "1002.03": ("港币户", "1002", 1, "HKD"),
    "1122.01": ("应收账款-国内", "1122", 1, "CNY"),
    "1122.02": ("应收账款-美元客户", "1122", 1, "USD"),
    "1122.03": ("应收账款-欧元客户", "1122", 1, "EUR"),
    "1123.01": ("预付款项-国内", "1123", 1, "CNY"),
    "1123.02": ("预付款项-日元", "1123", 1, "JPY"),
    "1123.03": ("预付款项-美元", "1123", 1, "USD"),
    "1405": ("库存商品", None, 1, "CNY"),
    "2202.01": ("应付账款-国内", "2202", -1, "CNY"),
    "2202.02": ("应付账款-美元供应商", "2202", -1, "USD"),
    "2202.03": ("应付账款-港币供应商", "2202", -1, "HKD"),
    "2203.01": ("预收款项-人民币", "2203", -1, "CNY"),
    "2203.02": ("预收款项-美元", "2203", -1, "USD"),
    "2221.01": ("应交增值税-进项税额", "2221", 1, "CNY"),
    "2221.02": ("应交增值税-销项税额", "2221", -1, "CNY"),
    "4001": ("实收资本", None, -1, "CNY"),
    "4103": ("本年利润", None, -1, "CNY"),
    "4104": ("利润分配-未分配利润", None, -1, "CNY"),
    "6001.01": ("出口收入(美元)", "6001", -1, "USD"),
    "6001.02": ("内销收入(人民币)", "6001", -1, "CNY"),
    "6401": ("主营业务成本", None, 1, "CNY"),
    "6601.01": ("运杂费(人民币)", "6601", 1, "CNY"),
    "6601.02": ("运杂费(港币)", "6601", 1, "HKD"),
    "6602.01": ("管理费用-办公费", "6602", 1, "CNY"),
    "6602.02": ("管理费用-差旅费", "6602", 1, "CNY"),
    "6602.03": ("管理费用-职工薪酬", "6602", 1, "CNY"),
    "6603.01": ("财务费用-手续费", "6603", 1, "CNY"),
    "6603.02": ("财务费用-汇兑损益", "6603", 1, "CNY"),
}
PARENTS = {
    "1002": "银行存款",
    "1122": "应收账款",
    "1123": "预付款项",
    "2202": "应付账款",
    "2203": "预收款项",
    "2221": "应交税费",
    "6001": "主营业务收入",
    "6601": "销售费用",
    "6602": "管理费用",
    "6603": "财务费用",
}
PL_ACCOUNTS = ["6001.01", "6001.02", "6401", "6601.01", "6601.02",
               "6602.01", "6602.02", "6602.03", "6603.01", "6603.02"]

# 期初余额(原币, 2026-01-01), 未列示科目期初为 0
OPEN = {
    "1002.01": 3500000.00,
    "1002.02": 200000.00,
    "1002.03": 500000.00,
    "1122.01": 860000.00,
    "1122.02": 320000.00,
    "1122.03": 150000.00,
    "1123.01": 120000.00,
    "1123.02": 5000000.00,
    "1405": 2600000.00,
    "2202.01": 640000.00,
    "2202.02": 180000.00,
    "2202.03": 300000.00,
    "2203.01": 150000.00,
    "2203.02": 40000.00,
    "4001": 8000000.00,
    "4104": 2023500.00,
}


def r2(x):
    return round(x + 0.0, 2)


# ---------------------------------------------------------------- 凭证引擎
VOUCHERS = []          # 每项: [日期, 摘要, [(科目, 币种, 借1/贷-1, 原币金额), ...]]
MONTH_NET = {}         # (月, 科目) -> [原币净额(借+), 本位币净额(借+)]


def line_cny(acct, ccy, orig):
    if ccy == "CNY":
        return r2(orig)
    return r2(orig * BOOK[ccy])


def post(d, summ, *lines):
    deb = sum(line_cny(l[0], l[1], l[3]) for l in lines if l[2] == 1)
    cre = sum(line_cny(l[0], l[1], l[3]) for l in lines if l[2] == -1)
    assert abs(deb - cre) < 0.005, "凭证本位币不平衡: %s %s 借%.2f 贷%.2f" % (d, summ, deb, cre)
    VOUCHERS.append([d, summ, list(lines)])
    for acct, ccy, dc, orig in lines:
        key = (d.month, acct)
        cny = line_cny(acct, ccy, orig)
        cur = MONTH_NET.get(key)
        if cur is None:
            cur = [0.0, 0.0]
            MONTH_NET[key] = cur
        cur[0] = r2(cur[0] + dc * (orig if ccy != "CNY" else 0.0))
        cur[1] = r2(cur[1] + dc * cny)


def fx_to_cny(d, summ, fx_acct, ccy, orig, bank_rate):
    """卖出外币收人民币(结汇/收外币直接结汇): 差价计入 6603.02"""
    bank = r2(orig * bank_rate)
    book = r2(orig * BOOK[ccy])
    diff = r2(bank - book)
    lines = [("1002.01", "CNY", 1, bank), (fx_acct, ccy, -1, orig)]
    if abs(diff) >= 0.005:
        lines.append(("6603.02", "CNY", (-1 if diff > 0 else 1), abs(diff)))
    post(d, summ, *lines)


def fx_from_cny(d, summ, fx_acct, ccy, orig, bank_rate):
    """以人民币购汇/购汇直接支付: 借外币科目, 贷1002.01, 价差计入 6603.02"""
    bank = r2(orig * bank_rate)
    book = r2(orig * BOOK[ccy])
    diff = r2(bank - book)
    lines = [(fx_acct, ccy, 1, orig), ("1002.01", "CNY", -1, bank)]
    if abs(diff) >= 0.005:
        lines.append(("6603.02", "CNY", (1 if diff > 0 else -1), abs(diff)))
    post(d, summ, *lines)


# ---------------------------------------------------------------- 月度业务数据
SALES_A = [120000, 80000, 120000, 140000, 150000, 160000]   # 出口销售-美国客户A (USD)
SALES_B = [80000, 60000, 80000, 80000, 90000, 100000]      # 出口销售-美国客户B (USD)
COLLECT = [100000, 100000, 120000, 160000, 160000, 180000]  # 收美国客户货款 (USD)
IMPORT_USD = [60000, 40000, 60000, 60000, 70000, 70000]     # 进口采购(美元供应商)
PAY_USD = [50000, 30000, 50000, 50000, 60000, 60000]        # 付美元供应商
DOM_BUY = [300000, 260000, 300000, 320000, 340000, 340000]  # 国内采购(不含税, 13%进项)
HKD_FREIGHT = [40000, 30000, 40000, 40000, 0, 45000]        # 港币运费(计提并支付)
SETTLE_USD = [60000, 40000, 60000, 60000, 70000, 70000]     # 结汇卖出美元
SETTLE_RATE = [7.15, 7.14, 7.15, 7.16, 7.15, 7.17]         # 结汇银行买入价
BUY_USD = [30000, 0, 30000, 30000, 30000, 30000]            # 购汇买入美元(2月春节资金安排不购汇)
BUY_RATE = [7.14, 7.13, 7.14, 7.15, 7.14, 7.15]            # 购汇银行卖出价
SALARY = 60000.00
OFFICE = 8000.00
TRAVEL = 15000.00
BANK_FEE = [1200, 900, 1200, 1500, 1200, 1500]
CARRY_COST = [800000, 700000, 850000, 850000, 900000, 900000]  # 结转销售成本
DOM_SALE = [200000, None, 200000, None, 200000, None]       # 内销收入(货款次月起收)
VAT = 0.13


def build_month(m):
    d = lambda day: date(2026, m, day)
    month_end = date(2026, m, calendar.monthrange(2026, m)[1])

    # 5日 出口销售
    post(d(5), "出口销售-美国客户A",
         ("1122.02", "USD", 1, float(SALES_A[m - 1])),
         ("6001.01", "USD", -1, float(SALES_A[m - 1])))
    if not (m in (2, 5)):  # 春节淡季/盘点月 B 客户无出运
        post(d(6), "出口销售-美国客户B",
             ("1122.02", "USD", 1, float(SALES_B[m - 1])),
             ("6001.01", "USD", -1, float(SALES_B[m - 1])))

    # 8日 (2月) 收到美国客户B预收货款
    if m == 2:
        post(d(8), "收到美国客户B预收货款",
             ("1002.02", "USD", 1, 20000.00),
             ("2203.02", "USD", -1, 20000.00))
    # 9日 (2月) 预付美元货款
    if m == 2:
        post(d(9), "预付美元供应商货款",
             ("1123.03", "USD", 1, 10000.00),
             ("1002.02", "USD", -1, 10000.00))
    # 10日 内销 / 收内销货款(6月货款于7月收回, 当月不体现)
    if DOM_SALE[m - 1]:
        amt = float(DOM_SALE[m - 1])
        post(d(10), "内销商品收入(含税%s)" % format(r2(amt * (1 + VAT)), ",.2f"),
             ("1122.01", "CNY", 1, r2(amt * (1 + VAT))),
             ("6001.02", "CNY", -1, amt),
             ("2221.02", "CNY", -1, r2(amt * VAT)))
    elif m in (2, 4):
        post(d(10), "收到内销客户货款",
             ("1002.01", "CNY", 1, 200000.00),
             ("1122.01", "CNY", -1, 200000.00))
    # 11日 (4月) 美元预付到货
    if m == 4:
        post(d(11), "美元预付货款到货入库",
             ("1405", "CNY", 1, 71200.00),
             ("1123.03", "USD", -1, 10000.00))
    # 12日 收美国客户货款
    post(d(12), "收到美国客户A货款",
         ("1002.02", "USD", 1, float(COLLECT[m - 1])),
         ("1122.02", "USD", -1, float(COLLECT[m - 1])))
    # 13日 (3月/6月) 购汇日元预付日本供应商
    if m in (3, 6):
        fx_from_cny(d(13), "购汇日元预付日本供应商货款", "1123.02", "JPY", 2000000.0, 0.0488)
    # 14日 工资
    post(d(14), "发放本月职工薪酬",
         ("6602.03", "CNY", 1, SALARY),
         ("1002.01", "CNY", -1, SALARY))
    # 15日 办公费/差旅费
    post(d(15), "报销办公费及差旅费",
         ("6602.01", "CNY", 1, OFFICE),
         ("6602.02", "CNY", 1, TRAVEL),
         ("1002.01", "CNY", -1, r2(OFFICE + TRAVEL)))
    # 16日 银行手续费
    post(d(16), "支付银行手续费",
         ("6603.01", "CNY", 1, float(BANK_FEE[m - 1])),
         ("1002.01", "CNY", -1, float(BANK_FEE[m - 1])))
    # 17日 (2月/4月/6月) 收欧元客户货款直接结汇
    if m == 2:
        fx_to_cny(d(17), "收到欧元客户货款并结汇", "1122.03", "EUR", 25000.0, 7.85)
    elif m == 4:
        fx_to_cny(d(17), "收到欧元客户货款并结汇", "1122.03", "EUR", 25000.0, 7.87)
    elif m == 6:
        fx_to_cny(d(17), "收到欧元客户货款并结汇", "1122.03", "EUR", 25000.0, 7.86)
    # 18日 进口采购
    imp = float(IMPORT_USD[m - 1])
    post(d(18), "进口采购-美元供应商",
         ("1405", "CNY", 1, r2(imp * BOOK["USD"])),
         ("2202.02", "USD", -1, imp))
    # 19日 (3月/6月) 预收美元货款转收入
    if m in (3, 6):
        post(d(19), "预收美元货款转出口收入",
             ("2203.02", "USD", 1, 20000.00),
             ("6001.01", "USD", -1, 20000.00))
    # 20日 国内采购
    dom = float(DOM_BUY[m - 1])
    post(d(20), "国内采购商品入库",
         ("1405", "CNY", 1, dom),
         ("2221.01", "CNY", 1, r2(dom * VAT)),
         ("2202.01", "CNY", -1, r2(dom * (1 + VAT))))
    # 22日 付美元供应商
    pay = float(PAY_USD[m - 1])
    post(d(22), "支付美元供应商货款",
         ("2202.02", "USD", 1, pay),
         ("1002.02", "USD", -1, pay))
    # 23/24日 计提并支付港币运费(5月货代合同间歇, 无运费)
    fr = float(HKD_FREIGHT[m - 1])
    if fr > 0:
        post(d(23), "计提香港货代运费",
             ("6601.02", "HKD", 1, fr),
             ("2202.03", "HKD", -1, fr))
        # 24日 支付港币运费
        post(d(24), "支付香港货代运费",
             ("2202.03", "HKD", 1, fr),
             ("1002.03", "HKD", -1, fr))
    # 25日 美元结汇
    fx_to_cny(d(25), "美元结汇", "1002.02", "USD",
              float(SETTLE_USD[m - 1]), SETTLE_RATE[m - 1])
    # 26日 购汇美元
    if BUY_USD[m - 1] > 0:
        fx_from_cny(d(26), "购汇美元", "1002.02", "USD",
                    float(BUY_USD[m - 1]), BUY_RATE[m - 1])
    # 26日 (5月) 购汇港币
    if m == 5:
        fx_from_cny(d(26), "购汇港币", "1002.03", "HKD", 80000.0, 0.919)
    # 27日 付国内供应商
    post(d(27), "支付本月国内供应商货款",
         ("2202.01", "CNY", 1, r2(dom * (1 + VAT))),
         ("1002.01", "CNY", -1, r2(dom * (1 + VAT))))
    # 28日 结转销售成本
    post(d(28), "结转本月销售成本",
         ("6401", "CNY", 1, float(CARRY_COST[m - 1])),
         ("1405", "CNY", -1, float(CARRY_COST[m - 1])))
    # 28日 (2月/5月) 人民币运杂费
    if m in (2, 5):
        post(d(28), "支付国内运杂费",
             ("6601.01", "CNY", 1, 6000.00),
             ("1002.01", "CNY", -1, 6000.00))
    # 月末 损益结转
    build_closing(month_end)


def build_closing(d):
    lines = []
    deb_total = 0.0
    cre_total = 0.0
    for acct in PL_ACCOUNTS:
        net = MONTH_NET.get((d.month, acct))
        if net is None or abs(net[1]) < 0.005:
            continue
        ccy = LEAVES[acct][3]
        orig = abs(net[1]) if ccy == "CNY" else abs(net[0])
        if net[1] > 0:  # 借方净额, 结转贷方
            lines.append((acct, ccy, -1, orig))
            cre_total = r2(cre_total + abs(net[1]))
        else:           # 贷方净额, 结转借方
            lines.append((acct, ccy, 1, orig))
            deb_total = r2(deb_total + abs(net[1]))
    plug = r2(deb_total - cre_total)
    if abs(plug) >= 0.005:
        if plug > 0:
            lines.append(("4103", "CNY", -1, plug))
        else:
            lines.append(("4103", "CNY", 1, -plug))
    post(d, "月末损益结转", *lines)


# ---------------------------------------------------------------- 序时账 (JE)
def build_journal():
    vouchers = sorted(VOUCHERS, key=lambda v: v[0])  # 稳定排序, 同日按建账顺序
    rows = []
    balances = {a: 0.0 for a in LEAVES}
    for a, (name, p, nat, ccy) in LEAVES.items():
        op = OPEN.get(a, 0.0)
        balances[a] = r2(op * nat if ccy != "CNY" else op * nat)
        # 期初有额科目按借贷性质定符号(借+/贷-), 余额列展示绝对值
    for idx, (d, summ, lines) in enumerate(vouchers, start=1):
        vno = "记-%04d" % idx
        for acct, ccy, dc, orig in lines:
            cny = line_cny(acct, ccy, orig)
            balances[acct] = r2(balances[acct] + dc * cny)
            bal = balances[acct]
            direction = "借" if bal > 0.004 else ("贷" if bal < -0.004 else "平")
            rows.append({
                "date": d, "vno": vno, "summary": summ,
                "acct": acct, "name": LEAVES[acct][0],
                "ccy": CNAME[ccy], "rate": BOOK[ccy] if ccy != "CNY" else None,
                "b_orig": orig if dc == 1 and ccy != "CNY" else None,
                "b_cny": cny if dc == 1 else None,
                "c_orig": orig if dc == -1 and ccy != "CNY" else None,
                "c_cny": cny if dc == -1 else None,
                "dir": direction, "bal": abs(bal) if abs(bal) > 0.004 else 0.0,
            })
    return rows


ANOMALY_SPECS = [
    # (谓词: 月份/科目/借贷/摘要包含, 替换后的币种写法)
    (lambda r: r["date"].month == 2 and r["acct"] == "1002.02" and r["b_orig"]
     and "收到美国客户A货款" in r["summary"], "usd"),
    (lambda r: r["date"].month == 4 and r["acct"] == "1002.02" and r["c_orig"]
     and "结汇" in r["summary"], "Usd"),
    (lambda r: r["date"].month == 6 and r["acct"] == "1122.02" and r["c_orig"]
     and "收到美国客户A货款" in r["summary"], "usd"),
]


def inject_anomalies(rows):
    hits = []
    for pred, repl in ANOMALY_SPECS:
        matched = [r for r in rows if pred(r)]
        assert len(matched) == 1, "异常币种注入定位失败: 命中 %d 行" % len(matched)
        matched[0]["ccy"] = repl
        hits.append(matched[0])
    return hits


# ---------------------------------------------------------------- 科目余额表 (TB)
def build_tb(je_rows):
    agg = {a: [0.0, 0.0, 0.0, 0.0] for a in LEAVES}  # 借原币/借本位币/贷原币/贷本位币
    for r in je_rows:
        a = r["acct"]
        if r["b_orig"]:
            agg[a][0] = r2(agg[a][0] + r["b_orig"])
        if r["b_cny"]:
            agg[a][1] = r2(agg[a][1] + r["b_cny"])
        if r["c_orig"]:
            agg[a][2] = r2(agg[a][2] + r["c_orig"])
        if r["c_cny"]:
            agg[a][3] = r2(agg[a][3] + r["c_cny"])

    leaf_rows = {}
    for a, (name, parent, nat, ccy) in LEAVES.items():
        op = OPEN.get(a, 0.0)
        op_s = r2(op * nat)
        op_cny = r2(op * BOOK[ccy]) if ccy != "CNY" else op
        bo, bc, co, cc = agg[a]
        end_s = r2(op_s + bo - co)
        end_sc = r2((op_cny * nat) + bc - cc)
        assert end_s * nat >= -0.005, "科目 %s 期末余额方向翻转" % a
        leaf_rows[a] = {
            "name": name, "ccy": ccy, "open_orig": op if ccy != "CNY" else None,
            "open_cny": op_cny,
            "bor_o": bo if ccy != "CNY" else None, "bor_c": bc,
            "cre_o": co if ccy != "CNY" else None, "cre_c": cc,
            "end_o": abs(end_s) if ccy != "CNY" else None,
            "end_c": abs(end_sc),
            "open_dir": ("借" if op_s > 0.004 else ("贷" if op_s < -0.004 else "平")),
            "end_dir": ("借" if end_sc > 0.004 else ("贷" if end_sc < -0.004 else "平")),
        }

    def sort_key(code):
        return tuple(int(p) for p in code.split("."))

    out = []
    for code in sorted(set(list(LEAVES.keys()) + list(PARENTS.keys())), key=sort_key):
        if code in PARENTS:
            kids = [leaf_rows[a] for a in LEAVES if LEAVES[a][1] == code]
            o = r2(sum(k["open_cny"] * (1 if k["open_dir"] == "借" else -1 if k["open_dir"] == "贷" else 0) for k in kids))
            b = r2(sum(k["bor_c"] for k in kids))
            c = r2(sum(k["cre_c"] for k in kids))
            e = r2(sum(k["end_c"] * (1 if k["end_dir"] == "借" else -1 if k["end_dir"] == "贷" else 0) for k in kids))
            out.append({"code": code, "name": PARENTS[code], "ccy": None,
                        "open_orig": None, "open_cny": abs(o), "bor_o": None, "bor_c": b,
                        "cre_o": None, "cre_c": c, "end_o": None, "end_c": abs(e),
                        "open_dir": ("借" if o > 0.004 else ("贷" if o < -0.004 else "平")),
                        "end_dir": ("借" if e > 0.004 else ("贷" if e < -0.004 else "平"))})
        else:
            row = leaf_rows[code]
            row["code"] = code
            out.append(row)

    # 合计行(仅本位币, 原币/币种留空)
    def side(v, dr):
        return v if dr == "借" else -v if dr == "贷" else 0.0

    open_side = r2(sum(side(l["open_cny"], l["open_dir"]) for l in leaf_rows.values()))
    bor = r2(sum(l["bor_c"] for l in leaf_rows.values()))
    cre = r2(sum(l["cre_c"] for l in leaf_rows.values()))
    end_side = r2(sum(side(l["end_c"], l["end_dir"]) for l in leaf_rows.values()))
    assert abs(open_side) < 0.01, "期初试算不平衡"
    assert abs(bor - cre) < 0.01, "本期发生试算不平衡"
    assert abs(end_side) < 0.01, "期末试算不平衡"
    out.append({"code": "合计", "name": "合计", "ccy": None,
                "open_orig": None, "open_cny": r2(sum(l["open_cny"] for l in leaf_rows.values()
                                                     if l["open_dir"] == "借")),
                "bor_o": None, "bor_c": bor, "cre_o": None, "cre_c": cre,
                "end_o": None, "end_c": r2(sum(l["end_c"] for l in leaf_rows.values()
                                               if l["end_dir"] == "借")),
                "open_dir": "平", "end_dir": "平"})
    return out, leaf_rows


# ---------------------------------------------------------------- 写 Excel
F_BODY = Font(name="宋体", size=10)
F_HEAD = Font(name="宋体", size=10, bold=True)
F_TITLE = Font(name="宋体", size=14, bold=True)
A_L = Alignment(horizontal="left", vertical="center")
A_C = Alignment(horizontal="center", vertical="center")
A_R = Alignment(horizontal="right", vertical="center")
AMT_FMT = "#,##0.00"
RATE_FMT = "0.00##"

TB_HEADERS = ["科目编码", "科目名称", "币种", "方向",
              "期初余额(原币)", "期初余额(本位币)",
              "本期发生借方(原币)", "本期发生借方(本位币)",
              "本期发生贷方(原币)", "本期发生贷方(本位币)",
              "期末余额(原币)", "期末余额(本位币)"]
JE_HEADERS = ["日期", "凭证字号", "摘要", "科目编码", "科目名称", "币种", "汇率",
              "借方原币", "借方本位币", "贷方原币", "贷方本位币", "余额方向", "余额"]


def put_amount(ws, row, col, value):
    if value is None:
        return
    cell = ws.cell(row=row, column=col, value=float(value))
    cell.number_format = AMT_FMT
    cell.font = F_BODY
    cell.alignment = A_R


def write_tb(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "科目余额表"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(TB_HEADERS))
    t = ws.cell(row=1, column=1, value="科目余额表")
    t.font = F_TITLE
    t.alignment = A_C
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(TB_HEADERS))
    info = ws.cell(row=2, column=1, value="单位：%s 2026年1月-6月 单位：元" % ENTITY)
    info.font = F_BODY
    info.alignment = A_L
    for j, h in enumerate(TB_HEADERS, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = F_HEAD
        c.alignment = A_C
    for i, row in enumerate(rows, start=4):
        for j, v in ((1, row["code"]), (2, row["name"]),
                     (3, CNAME.get(row["ccy"]) if row["ccy"] else None),
                     (4, row["end_dir"])):
            if v is not None:
                c = ws.cell(row=i, column=j, value=v)
                c.font = F_BODY
                c.alignment = A_L if j == 2 else A_C
        put_amount(ws, i, 5, row["open_orig"])
        put_amount(ws, i, 6, row["open_cny"])
        put_amount(ws, i, 7, row["bor_o"])
        put_amount(ws, i, 8, row["bor_c"])
        put_amount(ws, i, 9, row["cre_o"])
        put_amount(ws, i, 10, row["cre_c"])
        put_amount(ws, i, 11, row["end_o"])
        put_amount(ws, i, 12, row["end_c"])
    widths = [11, 22, 8, 6] + [15] * 8
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"
    wb.save(path)


def write_je(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "序时账"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(JE_HEADERS))
    t = ws.cell(row=1, column=1, value="序时账")
    t.font = F_TITLE
    t.alignment = A_C
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(JE_HEADERS))
    info = ws.cell(row=2, column=1,
                   value="单位：%s 期间：2026年1月1日至2026年6月30日 单位：元" % ENTITY)
    info.font = F_BODY
    info.alignment = A_L
    for j, h in enumerate(JE_HEADERS, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = F_HEAD
        c.alignment = A_C
    for i, r in enumerate(rows, start=4):
        c = ws.cell(row=i, column=1, value=r["date"])
        c.number_format = "yyyy-mm-dd"
        c.font = F_BODY
        c.alignment = A_C
        for j, v in ((2, r["vno"]), (3, r["summary"]), (4, r["acct"]), (5, r["name"]),
                     (6, r["ccy"]), (12, r["dir"])):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F_BODY
            c.alignment = A_L if j in (3, 5) else A_C
        if r["rate"] is not None:
            c = ws.cell(row=i, column=7, value=float(r["rate"]))
            c.number_format = RATE_FMT
            c.font = F_BODY
            c.alignment = A_C
        put_amount(ws, i, 8, r["b_orig"])
        put_amount(ws, i, 9, r["b_cny"])
        put_amount(ws, i, 10, r["c_orig"])
        put_amount(ws, i, 11, r["c_cny"])
        put_amount(ws, i, 13, r["bal"])
    widths = [11, 10, 26, 11, 20, 8, 8, 13, 14, 13, 14, 9, 14]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"
    wb.save(path)


# ---------------------------------------------------------------- 自检
CHECKS = []


def check(name, ok, detail=""):
    CHECKS.append((name, bool(ok), detail))
    print("[%s] %s%s" % ("通过" if ok else "失败", name, (" -- " + detail) if detail else ""))


def find_header(ws, keys):
    for row in ws.iter_rows(min_row=1, max_row=8):
        values = [str(c.value) if c.value is not None else "" for c in row]
        if all(any(k in v for v in values) for k in keys):
            return row[0].row, values
    return None, None


def self_check(tb_path, je_path, expect_anomalies):
    tb = load_workbook(tb_path).active
    je = load_workbook(je_path).active
    tb_hrow, tb_headers = find_header(tb, ["科目编码", "币种"])
    je_hrow, je_headers = find_header(je, ["凭证字号", "币种"])
    check("TB 表头行可定位(第%s行)且包含币种列" % tb_hrow, tb_hrow == 3)
    check("JE 表头行可定位(第%s行)且包含币种列" % je_hrow, je_hrow == 3)

    def col_of(headers, key):
        for i, h in enumerate(headers):
            if key in h:
                return i
        return None

    tb_rows = [[c.value for c in r] for r in tb.iter_rows(min_row=tb_hrow + 1)]
    tb_rows = [r for r in tb_rows if any(v is not None for v in r)]
    je_rows = [[c.value for c in r] for r in je.iter_rows(min_row=je_hrow + 1)]
    je_rows = [r for r in je_rows if any(v is not None for v in r)]

    # 1. 必需列存在
    tb_need = ["币种", "原币", "本位币", "借方", "贷方", "期初", "期末"]
    je_need = ["币种", "原币", "本位币", "借方", "贷方"]
    check("TB 列组完整(币种/原币/本位币/借贷/期初/期末)",
          all(any(k in h for h in tb_headers) for k in tb_need)
          and sum(1 for h in tb_headers if "原币" in h) >= 4
          and sum(1 for h in tb_headers if "本位币" in h) >= 4)
    check("JE 列组完整(币种/借方原币/借方本位币/贷方原币/贷方本位币)",
          all(any(k in h for h in je_headers) for k in je_need)
          and any("借方原币" in h for h in je_headers)
          and any("贷方本位币" in h for h in je_headers))

    ix = {k: col_of(tb_headers, k) for k in
          ["科目编码", "币种", "方向", "期初余额(原币)", "期初余额(本位币)",
           "本期发生借方(原币)", "本期发生借方(本位币)", "本期发生贷方(原币)",
           "本期发生贷方(本位币)", "期末余额(原币)", "期末余额(本位币)"]}
    jx = {k: col_of(je_headers, k) for k in
          ["日期", "凭证字号", "科目编码", "币种", "借方原币", "借方本位币",
           "贷方原币", "贷方本位币"]}
    check("TB/JE 关键列下标齐全", all(v is not None for v in list(ix.values()) + list(jx.values())))

    # 2. 列非全空 + 值域
    tb_ccy_vals = [str(r[ix["币种"]]) for r in tb_rows if r[ix["币种"]] is not None]
    fx_tb = sum(1 for r in tb_rows
                if any(r[ix[k]] is not None for k in ["期初余额(原币)", "本期发生借方(原币)",
                                                      "本期发生贷方(原币)", "期末余额(原币)"]))
    check("TB 币种列非全空(%d 行有币种)" % len(tb_ccy_vals), len(tb_ccy_vals) >= 20)
    check("TB 原币列组非全空(%d 行含原币金额)" % fx_tb, fx_tb >= 10)
    check("TB 本位币列非全空", sum(1 for r in tb_rows if r[ix["期末余额(本位币)"]] is not None) >= 30)
    check("TB 币种取值合法", set(tb_ccy_vals) <= set(CNAME.values()))

    je_ccy_vals = [str(r[jx["币种"]]) for r in je_rows if r[jx["币种"]] is not None]
    anomalies = [v for v in je_ccy_vals if v.lower() == "usd" and v != "美元"]
    normal_ok = set(v for v in je_ccy_vals if v not in ("usd", "Usd")) <= set(CNAME.values())
    check("JE 币种列非全空(%d 行)" % len(je_ccy_vals), len(je_ccy_vals) >= len(je_rows) - 5)
    check("JE 币种常规取值合法(异常写法除外)", normal_ok)
    check("JE 币种异常写法恰好 %d 行(当前 %d 行: %s)"
          % (len(expect_anomalies), len(anomalies), "/".join(anomalies)),
          len(anomalies) == len(expect_anomalies))
    check("JE 原币列非全空", sum(1 for r in je_rows
                              if r[jx["借方原币"]] is not None or r[jx["贷方原币"]] is not None) >= 50)
    check("JE 本位币列非全空", sum(1 for r in je_rows
                               if r[jx["借方本位币"]] is not None or r[jx["贷方本位币"]] is not None) >= 100)

    # 3. 数值/日期类型
    amt_bad, date_bad = [], []
    for r in je_rows:
        if not isinstance(r[jx["日期"]], (datetime, date)):
            date_bad.append(r[jx["日期"]])
        for k in ["借方原币", "借方本位币", "贷方原币", "贷方本位币"]:
            v = r[jx[k]]
            if v is not None and (isinstance(v, bool) or not isinstance(v, (int, float))):
                amt_bad.append((k, v))
    check("JE 日期列全部为日期型", not date_bad)
    check("JE 金额列全部为数值型(无文本/千分位字符串)", not amt_bad)
    tb_amt_bad = []
    for r in tb_rows:
        for k in ["期初余额(原币)", "期初余额(本位币)", "本期发生借方(原币)", "本期发生借方(本位币)",
                  "本期发生贷方(原币)", "本期发生贷方(本位币)", "期末余额(原币)", "期末余额(本位币)"]:
            v = r[ix[k]]
            if v is not None and (isinstance(v, bool) or not isinstance(v, (int, float))):
                tb_amt_bad.append((k, v))
    check("TB 金额列全部为数值型", not tb_amt_bad)

    # 4. 每凭证本位币借贷平衡
    vouchers = {}
    for r in je_rows:
        vno = str(r[jx["凭证字号"]])
        d = vouchers.setdefault(vno, [0.0, 0.0])
        d[0] += float(r[jx["借方本位币"]] or 0.0)
        d[1] += float(r[jx["贷方本位币"]] or 0.0)
    unbalanced = {k: v for k, v in vouchers.items() if abs(v[0] - v[1]) > 0.005}
    check("每张凭证本位币借贷平衡(%d 张凭证)" % len(vouchers), not unbalanced,
          "不平凭证: %s" % list(unbalanced)[:5] if unbalanced else "")

    # 5. TB 滚动勾稽 (叶子行: 期末 = 期初 + 借 - 贷, 原币+本位币)
    def signed(v, dr):
        s = 1 if dr == "借" else (-1 if dr == "贷" else 0)
        return float(v or 0.0) * s

    roll_bad = []
    for r in tb_rows:
        code = str(r[ix["科目编码"]])
        ccy = r[ix["币种"]]
        if code in ("合计",) or ccy is None:
            continue  # 父级/合计行仅有本位币汇总
        e = signed(r[ix["期末余额(本位币)"]], r[ix["方向"]])
        expect = signed(r[ix["期初余额(本位币)"]], "借" if r[ix["方向"]] == "借" else "贷") \
            if r[ix["期初余额(本位币)"]] else 0.0
        # 期初方向需要单独推断: 用借贷性质重算更稳 -> 直接用代数: 期末-借+贷 应为期初(符号未知),
        # 改为校验 |期末-(0+借-贷+期初带方向)| -- 为此按科目自然方向回推期初符号。
        o = float(r[ix["期初余额(本位币)"]] or 0.0)
        nat = LEAVES.get(code, (None, None, 1, None))[2]
        o_s = o * nat
        if abs((o_s + float(r[ix["本期发生借方(本位币)"]] or 0.0)
                - float(r[ix["本期发生贷方(本位币)"]] or 0.0)) - e) > 0.005:
            roll_bad.append((code, "本位币"))
        if ccy in ("美元", "港币", "欧元", "日元"):
            oo = float(r[ix["期初余额(原币)"]] or 0.0) * nat
            ee = float(r[ix["期末余额(原币)"]] or 0.0) * (1 if e >= 0 else -1)
            if abs((oo + float(r[ix["本期发生借方(原币)"]] or 0.0)
                    - float(r[ix["本期发生贷方(原币)"]] or 0.0)) - ee) > 0.005:
                roll_bad.append((code, "原币"))
    check("TB 叶子行滚动勾稽(期末=期初+借-贷, 原币+本位币, 除父级/合计行)",
          not roll_bad, str(roll_bad[:5]) if roll_bad else "")

    # 6. TB 试算平衡
    leaves_only = [r for r in tb_rows if str(r[ix["科目编码"]]) in LEAVES]
    open_ss = sum(float(r[ix["期初余额(本位币)"]] or 0.0) * LEAVES[str(r[ix["科目编码"]])][2]
                  for r in leaves_only)
    end_ss = sum(signed(r[ix["期末余额(本位币)"]], r[ix["方向"]]) for r in leaves_only)
    bor_t = sum(float(r[ix["本期发生借方(本位币)"]] or 0.0) for r in leaves_only)
    cre_t = sum(float(r[ix["本期发生贷方(本位币)"]] or 0.0) for r in leaves_only)
    check("TB 期初试算平衡(借=贷)", abs(open_ss) < 0.01, "差额 %.2f" % open_ss)
    check("TB 期末试算平衡(借=贷)", abs(end_ss) < 0.01, "差额 %.2f" % end_ss)
    check("TB 本期发生借方合计=贷方合计", abs(bor_t - cre_t) < 0.01,
          "借 %.2f / 贷 %.2f" % (bor_t, cre_t))

    # 7. 规模
    check("TB 数据行数在 40-80 (当前 %d)" % len(tb_rows), 40 <= len(tb_rows) <= 80)
    check("JE 数据行数在 150-300 (当前 %d)" % len(je_rows), 150 <= len(je_rows) <= 300)

    # 8. 预埋点1: 期末未重估
    hist_ok = True
    detail = []
    for r in tb_rows:
        code = str(r[ix["科目编码"]])
        if code in ("1002.02", "1122.02", "1122.03", "1002.03"):
            eo = float(r[ix["期末余额(原币)"]] or 0.0)
            ec = float(r[ix["期末余额(本位币)"]] or 0.0)
            ccy = {"1002.02": "USD", "1122.02": "USD", "1122.03": "EUR", "1002.03": "HKD"}[code]
            if abs(ec - r2(eo * BOOK[ccy])) > 0.005:
                hist_ok = False
                detail.append(code)
    check("预埋点1前置: 1002.02/1122.02/1122.03/1002.03 期末本位币=期末原币x记账汇率(按历史汇率滚出)",
          hist_ok, str(detail))
    reval_rows = [r for r in je_rows
                  if str(r[jx["科目编码"]]) == "6603.02"
                  and any(w in str(r[col_of(je_headers, '摘要')] or "")
                          for w in ("重估", "期末汇率", "汇兑调整"))]
    check("预埋点1: JE 无任何期末汇率重估分录(6603.02 无重估摘要)",
          not reval_rows)
    jun30_fx = [r for r in je_rows
                if isinstance(r[jx["日期"]], (datetime, date))
                and r[jx["日期"]].month == 6 and r[jx["日期"]].day == 30
                and str(r[jx["科目编码"]]) in ("1002.02", "1122.02")]
    check("预埋点1: 6月30日外币货币性科目无任何发生额(仅有损益结转)",
          not jun30_fx)
    return tb_rows, je_rows


def print_exposure(leaf_rows):
    print()
    print("=" * 62)
    print("预埋测试点 1 -- 期末未重估 应重估金额测算 (工具应提示)")
    print("=" * 62)
    def endo(code):
        return leaf_rows[code]["end_o"] or 0.0, leaf_rows[code]["end_dir"]

    lines = [
        ("USD", [("1002.02", 1), ("1122.02", 1), ("2202.02", -1)], True),
        ("EUR", [("1122.03", 1)], True),
        ("HKD", [("1002.03", 1), ("2202.03", -1)], True),
        ("JPY", [("1123.02", 1)], False),  # 预付款项为非货币性项目(参考提示)
    ]
    total = 0.0
    for ccy, items, monetary in lines:
        net = 0.0
        parts = []
        for code, sign in items:
            o, d = endo(code)
            net += sign * o
            parts.append("%s%s %s" % (LEAVES[code][0], "借" if sign > 0 else "(-)", format(o, ",.2f")))
        move = END_RATE[ccy] - BOOK[ccy]
        effect = r2(net * move)
        if monetary:
            total = r2(total + effect)
        print("  %s 净敞口 = %s" % (ccy, " + ".join(parts) if len(parts) == 1 else " + ".join(parts)))
        print("      净额 %s %s, 期末汇率%.4f - 记账汇率%.4f = %.4f, 应重估 %s%s"
              % (format(net, ",.2f"), CNAME[ccy], END_RATE[ccy], BOOK[ccy], move,
                 format(abs(effect), ",.2f"), "收益(漏提)" if effect > 0 else "损失(漏提)"))
        if not monetary:
            print("      (注: 预付/预收为非货币性项目, 严格按准则不重估, 工具可作参考提示)")
    print("  ---- 货币性项目合计应重估(全部为漏提的汇兑收益): %s 元" % format(total, ",.2f"))
    return total


def main():
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    for m in range(1, 7):
        build_month(m)

    je_rows = build_journal()
    anomalies = inject_anomalies(je_rows)
    tb_rows, leaf_rows = build_tb(je_rows)

    tb_path = os.path.join(OUT_DIR, "用友_科目余额表.xlsx")
    je_path = os.path.join(OUT_DIR, "用友_序时账.xlsx")
    write_tb(tb_path, tb_rows)
    write_je(je_path, je_rows)

    print("生成完成:")
    print("  %s (数据行 x%d)" % (tb_path, len(tb_rows)))
    print("  %s (数据行 x%d, 凭证 %d 张)" % (je_path, len(je_rows), len(set(r["vno"] for r in je_rows))))
    print()
    print("=" * 62)
    print("自检")
    print("=" * 62)
    self_check(tb_path, je_path, anomalies)
    total = print_exposure(leaf_rows)

    print()
    print("=" * 62)
    print("预埋测试点 2 -- 币种大小写混用 (位于序时账, 工具应归一化为 USD/美元)")
    print("=" * 62)
    for r in anomalies:
        excel_row = 4 + je_rows.index(r)
        print("  第%d行 %s %s 摘要[%s] 科目%s 币种列写作[%s] 原币%.2f"
              % (excel_row, r["date"].strftime("%Y-%m-%d"), r["vno"], r["summary"],
                 r["acct"], r["ccy"], (r["b_orig"] or r["c_orig"] or 0.0)))
    print()
    failed = [c for c in CHECKS if not c[1]]
    print("自检汇总: %d 项检查, 通过 %d 项, 失败 %d 项" % (len(CHECKS), len(CHECKS) - len(failed), len(failed)))
    if failed:
        for name, _, detail in failed:
            print("  失败: %s %s" % (name, detail))
        sys.exit(1)
    print("全部通过。")


if __name__ == "__main__":
    main()
