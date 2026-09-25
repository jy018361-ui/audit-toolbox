# -*- coding: utf-8 -*-
"""
金蝶(K/3 Cloud / KIS-WIN)导出风格 汇率损益测试数据集 生成 + 自检 一体脚本

虚构主体: 华金机械制造有限公司 (本位币 CNY, 期间 2026-01-01 至 2026-06-30, 第 1-6 期)
输出(与本脚本同目录):
    金蝶_科目余额表.xlsx  -- 金蝶科目余额表风格:
        第1行标题"科目余额表", 第2行 单位/期间/币别 导出信息,
        第3-4行两行复合表头: 期初余额|本期发生|期末余额 (上) x 借方/贷方 x 原币金额/本位币金额 (下)
    金蝶_序时账.xlsx      -- 金蝶核算项目明细账/序时账风格:
        日期/凭证字/凭证号/摘要/科目编码/科目名称/币种/借方(原币)/贷方(原币)/
        借方(本位币)/贷方(本位币)/余额(原币)/余额(本位币)

统一业务设定(四套 ERP 测试集共用):
    记账汇率: USD=7.12 EUR=7.83 HKD=0.91 JPY=0.0482
    期末汇率(2026-06-30): USD=7.18 EUR=7.90 HKD=0.915 JPY=0.0495
本套补充的虚构月末汇率(1-5月, 仅用于月末调汇凭证; 6月末即统一期末汇率):
    USD 7.08/7.05/7.15/7.10/7.20   EUR 7.80/7.76/7.88/7.85/7.92
    HKD 0.905/0.902/0.908/0.912/0.918   JPY 0.0478/0.0475/0.0485/0.0480/0.0490

预埋审计测试点(详见 README.md):
    [点1] 折算汇率异常: 3 张美元凭证(记-18 / 记-49 / 记-63, 共 6 个分录行)
          本位币按偏离汇率 6.80 折算(正确应按记账汇率 7.12), 凭证本身借贷仍平衡。
    [点2] 期初余额不勾稽: 1002.02 美元户 TB 期初原币 850,000.00(正确 800,000.00,
          虚增 50,000.00 USD = 356,000.00 CNY), 差额同步虚挂于 4104 未分配利润
          期初贷方以维持试算平衡; 两科目期末余额均按正确口径列示, 因此"期初+本期
          发生 != 期末"恰好在两个科目上分别暴露 -50,000/-356,000 与 +356,000。

基准之外的全部数据自洽: 每凭证本位币借贷平衡、其余科目 TB 滚动勾稽一致、
外币货币性项目期末余额 = 期末原币 x 期末汇率(月末调汇完整)。
控制台输出不使用 emoji (Windows GBK 兼容), 自检逐项打印, 任一失败退出码非 0。
复现: python generate.py    (依赖 openpyxl)
"""
import sys
from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent
ENTITY = "华金机械制造有限公司"
TB_PATH = OUT_DIR / "金蝶_科目余额表.xlsx"
JE_PATH = OUT_DIR / "金蝶_序时账.xlsx"

TWO = Decimal("0.01")
ZERO = Decimal("0")


def r2(x: Decimal) -> Decimal:
    return x.quantize(TWO, rounding=ROUND_HALF_UP)


def D(s) -> Decimal:
    return Decimal(str(s))


def fmt(d: Decimal) -> str:
    return "{:,.2f}".format(d)


CUR_CN = {"CNY": "人民币", "USD": "美元", "EUR": "欧元", "HKD": "港元", "JPY": "日元"}
CUR_ORDER = {"CNY": 0, "USD": 1, "EUR": 2, "HKD": 3, "JPY": 4}

# 记账汇率(统一设定) 与 期末汇率(2026-06-30, 统一设定)
BOOK = {"CNY": D("1"), "USD": D("7.12"), "EUR": D("7.83"),
        "HKD": D("0.91"), "JPY": D("0.0482")}
END_RATE = {"CNY": D("1"), "USD": D("7.18"), "EUR": D("7.90"),
            "HKD": D("0.915"), "JPY": D("0.0495")}
# 各月末汇率(虚构), 第 6 期即统一设定的期末汇率
MER = {
    1: {"USD": D("7.08"), "EUR": D("7.80"), "HKD": D("0.905"), "JPY": D("0.0478")},
    2: {"USD": D("7.05"), "EUR": D("7.76"), "HKD": D("0.902"), "JPY": D("0.0475")},
    3: {"USD": D("7.15"), "EUR": D("7.88"), "HKD": D("0.908"), "JPY": D("0.0485")},
    4: {"USD": D("7.10"), "EUR": D("7.85"), "HKD": D("0.912"), "JPY": D("0.0480")},
    5: {"USD": D("7.20"), "EUR": D("7.92"), "HKD": D("0.918"), "JPY": D("0.0490")},
    6: {"USD": D("7.18"), "EUR": D("7.90"), "HKD": D("0.915"), "JPY": D("0.0495")},
}
ANOM_RATE = D("6.80")

# 预埋点2: 1002.02 期初原币虚增 50,000.00 USD
OPEN_DIFF_USD = D("50000")
OPEN_DIFF_CNY = r2(OPEN_DIFF_USD * BOOK["USD"])  # 356,000.00

SUBJ_NAME = {
    "1001": "库存现金",
    "1002.01": "银行存款-人民币户",
    "1002.02": "银行存款-美元户",
    "1002.03": "银行存款-欧元户",
    "1002.04": "银行存款-港元户",
    "1002.05": "银行存款-日元户",
    "1012": "其他货币资金-信用证保证金(美元)",
    "1122.01": "应收账款-国内客户",
    "1122.02": "应收账款-美元客户",
    "1122.03": "应收账款-欧元客户",
    "1122.04": "应收账款-港元客户",
    "1123": "预付账款-国内供应商",
    "1221": "其他应收款-备用金",
    "1403.01": "原材料-钢材",
    "1403.02": "原材料-标准件",
    "1405": "库存商品",
    "1601": "固定资产",
    "1602": "累计折旧",
    "1604": "在建工程",
    "2202.01": "应付账款-国内供应商",
    "2202.02": "应付账款-美元供应商",
    "2202.03": "应付账款-日本供应商",
    "2211": "应付职工薪酬",
    "2221.01": "应交税费-应交增值税(销项税额)",
    "2221.02": "应交税费-应交增值税(进项税额)",
    "2241": "其他应付款",
    "2501": "长期借款",
    "4001": "实收资本",
    "4104": "利润分配-未分配利润",
    "6001.01": "主营业务收入-出口",
    "6001.02": "主营业务收入-境内",
    "6051": "其他业务收入",
    "6401": "主营业务成本",
    "6601.01": "销售费用-职工薪酬",
    "6601.02": "销售费用-运输费",
    "6602.01": "管理费用-职工薪酬",
    "6602.02": "管理费用-折旧费",
    "6602.03": "管理费用-办公费",
    "6603.01": "财务费用-利息支出",
    "6603.02": "财务费用-汇兑损益",
}

# 外币货币性项目(参与月末调汇)
MONETARY_FX = {"1002.02", "1002.03", "1002.04", "1002.05", "1012",
               "1122.02", "1122.03", "1122.04", "2202.02", "2202.03"}

# 期初余额(2026-01-01, 原币, 借方为正/贷方为负); 4104 为轧差平衡数, 由代码计算
OPEN_ORIG = {
    ("1001", "CNY"): D("45000"),
    ("1002.01", "CNY"): D("3280000"),
    ("1002.02", "USD"): D("800000"),      # 正确值; TB 中按预埋点虚增为 850,000
    ("1002.03", "EUR"): D("320000"),
    ("1002.04", "HKD"): D("250000"),
    ("1002.05", "JPY"): D("8000000"),
    ("1012", "USD"): D("120000"),
    ("1122.01", "CNY"): D("1450000"),
    ("1122.02", "USD"): D("260000"),
    ("1122.03", "EUR"): D("145000"),
    ("1122.04", "HKD"): D("80000"),
    ("1123", "CNY"): D("85000"),
    ("1221", "CNY"): D("30000"),
    ("1403.01", "CNY"): D("1120000"),
    ("1403.02", "CNY"): D("460000"),
    ("1405", "CNY"): D("2600000"),
    ("1601", "CNY"): D("6850000"),
    ("1602", "CNY"): D("-2140000"),
    ("2202.01", "CNY"): D("-980000"),
    ("2202.02", "USD"): D("-175000"),
    ("2202.03", "JPY"): D("-3000000"),
    ("2221.01", "CNY"): D("-65000"),
    ("2221.02", "CNY"): D("38000"),
    ("2241", "CNY"): D("-120000"),
    ("2501", "CNY"): D("-2000000"),
    ("4001", "CNY"): D("-8000000"),
}

# ---------------------------------------------------------------- 记账引擎
entries = []          # 序时账分录(有序)
voucher_no = 0
bal_orig = {}         # (科目, 币种) -> 原币余额(借正贷负, 滚动)
bal_base = {}
anomaly_rows = []     # 预埋点1: (凭证号, 日期, 科目, 摘要, 原币, 本位币, 隐含汇率)


def init_openings():
    for (s, c), o in OPEN_ORIG.items():
        bal_orig[(s, c)] = o
        bal_base[(s, c)] = o if c == "CNY" else r2(o * BOOK[c])
    plug = -sum(bal_base.values())          # 4104 轧差(正确口径, 贷方余额)
    bal_orig[("4104", "CNY")] = plug
    bal_base[("4104", "CNY")] = plug
    return plug


PLUG_4104 = init_openings()
OPEN_SNAP_O = dict(bal_orig)                # 期初快照(正确口径)
OPEN_SNAP_B = dict(bal_base)


def L(subj, cur, dr=None, cr=None, rate=None):
    """生成一行分录数据(元组). 默认本位币按记账汇率折算; 人民币行原币=本位币."""
    rt = rate if rate is not None else BOOK[cur]
    if cur == "CNY":
        return (subj, cur, dr, cr, dr, cr)
    bd = r2(dr * rt) if dr else None
    bc = r2(cr * rt) if cr else None
    return (subj, cur, dr, cr, bd, bc)


def V(d, summ, rows):
    """过账一张凭证: rows 为 L()/手工元组列表, 自动更新余额."""
    global voucher_no
    voucher_no += 1
    for (s, c, od, oc, bd, bc) in rows:
        entries.append({"date": d, "word": "记", "no": voucher_no, "summ": summ,
                        "subj": s, "cur": c, "od": od, "oc": oc, "bd": bd, "bc": bc})
        k = (s, c)
        bal_orig[k] = bal_orig.get(k, ZERO) + (od or ZERO) - (oc or ZERO)
        bal_base[k] = bal_base.get(k, ZERO) + (bd or ZERO) - (bc or ZERO)


def VA(d, summ, rows, anom_subj):
    """过账带预埋点1(偏离汇率 6.80)的凭证, 并登记异常行."""
    V(d, summ, rows)
    for (s, c, od, oc, bd, bc) in rows:
        if s in anom_subj and c != "CNY":
            orig = od or oc
            base = bd or bc
            anomaly_rows.append((voucher_no, d, s, summ, orig, base, r2(base / orig)))


def reval(m, d):
    """月末调汇: 外币货币性科目余额调整至月末汇率, 差额记 6603.02(原币为空)."""
    rows = []
    net = ZERO
    for k in sorted(bal_orig):
        s, c = k
        if c == "CNY" or s not in MONETARY_FX:
            continue
        o, b = bal_orig[k], bal_base[k]
        if o == 0 and b == 0:
            continue
        delta = r2(o * MER[m][c]) - b
        if delta == 0:
            continue
        if delta > 0:
            rows.append((s, c, None, None, delta, None))
        else:
            rows.append((s, c, None, None, None, -delta))
        net += delta
    if not rows:
        return
    if net > 0:
        rows.append(("6603.02", "CNY", None, net, None, net))
    else:
        rows.append(("6603.02", "CNY", -net, None, -net, None))
    V(d, "月末汇兑损益结转", rows)


# ---------------------------------------------------------------- 业务凭证
def build():
    # ---- 1 月 ----
    V(date(2026, 1, 5), "购汇 100,000 美元(购汇价 7.15)", [
        L("1002.02", "USD", dr=D("100000")),
        L("6603.02", "CNY", dr=D("3000")),
        L("1002.01", "CNY", cr=D("715000")),
    ])
    V(date(2026, 1, 8), "出口销售-美国 ACME 公司", [
        L("1122.02", "USD", dr=D("180000")),
        L("6001.01", "USD", cr=D("180000")),
    ])
    V(date(2026, 1, 12), "出口销售-德国 Kraft 公司", [
        L("1122.03", "EUR", dr=D("90000")),
        L("6001.01", "EUR", cr=D("90000")),
    ])
    V(date(2026, 1, 15), "收到 ACME 公司货款", [
        L("1002.02", "USD", dr=D("150000")),
        L("1122.02", "USD", cr=D("150000")),
    ])
    V(date(2026, 1, 18), "缴纳信用证保证金", [
        L("1012", "USD", dr=D("60000")),
        L("1002.02", "USD", cr=D("60000")),
    ])
    V(date(2026, 1, 20), "采购钢材(13% 增值税)", [
        L("1403.01", "CNY", dr=D("420000")),
        L("2221.02", "CNY", dr=D("54600")),
        L("2202.01", "CNY", cr=D("474600")),
    ])
    V(date(2026, 1, 22), "出口销售-香港中贸公司", [
        L("1122.04", "HKD", dr=D("500000")),
        L("6001.01", "HKD", cr=D("500000")),
    ])
    V(date(2026, 1, 25), "支付国内供应商货款", [
        L("2202.01", "CNY", dr=D("350000")),
        L("1002.01", "CNY", cr=D("350000")),
    ])
    V(date(2026, 1, 26), "采购日本标准件(赊购)", [
        L("1403.02", "CNY", dr=D("120500")),
        L("2202.03", "JPY", cr=D("2500000")),
    ])
    V(date(2026, 1, 27), "提取备用现金", [
        L("1001", "CNY", dr=D("30000")),
        L("1002.01", "CNY", cr=D("30000")),
    ])
    V(date(2026, 1, 28), "计提 1 月工资", [
        L("6601.01", "CNY", dr=D("68000")),
        L("6602.01", "CNY", dr=D("102000")),
        L("2211", "CNY", cr=D("170000")),
    ])
    V(date(2026, 1, 29), "发放 1 月工资", [
        L("2211", "CNY", dr=D("170000")),
        L("1002.01", "CNY", cr=D("170000")),
    ])
    V(date(2026, 1, 31), "计提折旧", [
        L("6602.02", "CNY", dr=D("28500")),
        L("1602", "CNY", cr=D("28500")),
    ])
    V(date(2026, 1, 31), "采购产成品(13% 增值税)", [
        L("1405", "CNY", dr=D("1690000")),
        L("2221.02", "CNY", dr=D("219700")),
        L("2202.01", "CNY", cr=D("1909700")),
    ])
    V(date(2026, 1, 31), "结转 1 月销售成本", [
        L("6401", "CNY", dr=D("1540000")),
        L("1405", "CNY", cr=D("1540000")),
    ])
    reval(1, date(2026, 1, 31))

    # ---- 2 月 ----
    V(date(2026, 2, 5), "结汇 200,000 美元(结汇价 7.06)", [
        L("1002.01", "CNY", dr=D("1412000")),
        L("6603.02", "CNY", dr=D("12000")),
        L("1002.02", "USD", cr=D("200000")),
    ])
    # [预埋点1] 本位币按 6.80 折算(正确应按 7.12, 差异 -38,400)
    VA(date(2026, 2, 9), "出口销售-美国 ACME 公司", [
        L("1122.02", "USD", dr=D("120000"), rate=ANOM_RATE),
        L("6001.01", "USD", cr=D("120000"), rate=ANOM_RATE),
    ], {"1122.02", "6001.01"})
    V(date(2026, 2, 12), "收到德国 Kraft 公司货款", [
        L("1002.03", "EUR", dr=D("60000")),
        L("1122.03", "EUR", cr=D("60000")),
    ])
    V(date(2026, 2, 16), "支付美元供应商 Norton 货款", [
        L("2202.02", "USD", dr=D("80000")),
        L("1002.02", "USD", cr=D("80000")),
    ])
    V(date(2026, 2, 20), "采购标准件(13% 增值税)", [
        L("1403.02", "CNY", dr=D("180000")),
        L("2221.02", "CNY", dr=D("23400")),
        L("2202.01", "CNY", cr=D("203400")),
    ])
    V(date(2026, 2, 24), "支付出口海运费", [
        L("6601.02", "CNY", dr=D("26000")),
        L("1002.01", "CNY", cr=D("26000")),
    ])
    V(date(2026, 2, 25), "收到香港中贸公司货款", [
        L("1002.04", "HKD", dr=D("300000")),
        L("1122.04", "HKD", cr=D("300000")),
    ])
    V(date(2026, 2, 27), "计提 2 月工资", [
        L("6601.01", "CNY", dr=D("68400")),
        L("6602.01", "CNY", dr=D("102600")),
        L("2211", "CNY", cr=D("171000")),
    ])
    V(date(2026, 2, 27), "发放 2 月工资", [
        L("2211", "CNY", dr=D("171000")),
        L("1002.01", "CNY", cr=D("171000")),
    ])
    V(date(2026, 2, 28), "计提折旧", [
        L("6602.02", "CNY", dr=D("28500")),
        L("1602", "CNY", cr=D("28500")),
    ])
    V(date(2026, 2, 28), "采购产成品(13% 增值税)", [
        L("1405", "CNY", dr=D("1050000")),
        L("2221.02", "CNY", dr=D("136500")),
        L("2202.01", "CNY", cr=D("1186500")),
    ])
    V(date(2026, 2, 28), "结转 2 月销售成本", [
        L("6401", "CNY", dr=D("1180000")),
        L("1405", "CNY", cr=D("1180000")),
    ])
    reval(2, date(2026, 2, 28))

    # ---- 3 月 ----
    V(date(2026, 3, 6), "出口销售-美国 ACME 公司", [
        L("1122.02", "USD", dr=D("210000")),
        L("6001.01", "USD", cr=D("210000")),
    ])
    V(date(2026, 3, 10), "收到 ACME 公司货款", [
        L("1002.02", "USD", dr=D("190000")),
        L("1122.02", "USD", cr=D("190000")),
    ])
    V(date(2026, 3, 11), "进口钢材(赊购, 美元结算)", [
        L("1403.01", "CNY", dr=D("1068000")),
        L("2202.02", "USD", cr=D("150000")),
    ])
    V(date(2026, 3, 13), "出口销售-意大利 MEP 公司", [
        L("1122.03", "EUR", dr=D("85000")),
        L("6001.01", "EUR", cr=D("85000")),
    ])
    V(date(2026, 3, 17), "购汇 120,000 美元(购汇价 7.16)", [
        L("1002.02", "USD", dr=D("120000")),
        L("6603.02", "CNY", dr=D("4800")),
        L("1002.01", "CNY", cr=D("859200")),
    ])
    V(date(2026, 3, 20), "支付美元供应商 Norton 货款", [
        L("2202.02", "USD", dr=D("100000")),
        L("1002.02", "USD", cr=D("100000")),
    ])
    V(date(2026, 3, 23), "报销办公费(备用金)", [
        L("6602.03", "CNY", dr=D("8600")),
        L("1221", "CNY", cr=D("8600")),
    ])
    V(date(2026, 3, 24), "支付日本供应商货款", [
        L("2202.03", "JPY", dr=D("2000000")),
        L("1002.05", "JPY", cr=D("2000000")),
    ])
    V(date(2026, 3, 26), "预付款到货转销", [
        L("1403.01", "CNY", dr=D("85000")),
        L("1123", "CNY", cr=D("85000")),
    ])
    V(date(2026, 3, 27), "支付一季度贷款利息", [
        L("6603.01", "CNY", dr=D("42000")),
        L("1002.01", "CNY", cr=D("42000")),
    ])
    V(date(2026, 3, 30), "购入设备(在建, 13% 增值税)", [
        L("1604", "CNY", dr=D("360000")),
        L("2221.02", "CNY", dr=D("46800")),
        L("1002.01", "CNY", cr=D("406800")),
    ])
    V(date(2026, 3, 31), "计提 3 月工资", [
        L("6601.01", "CNY", dr=D("68000")),
        L("6602.01", "CNY", dr=D("102000")),
        L("2211", "CNY", cr=D("170000")),
    ])
    V(date(2026, 3, 31), "发放 3 月工资", [
        L("2211", "CNY", dr=D("170000")),
        L("1002.01", "CNY", cr=D("170000")),
    ])
    V(date(2026, 3, 31), "计提折旧", [
        L("6602.02", "CNY", dr=D("28500")),
        L("1602", "CNY", cr=D("28500")),
    ])
    V(date(2026, 3, 31), "采购产成品(13% 增值税)", [
        L("1405", "CNY", dr=D("1800000")),
        L("2221.02", "CNY", dr=D("234000")),
        L("2202.01", "CNY", cr=D("2034000")),
    ])
    V(date(2026, 3, 31), "结转 3 月销售成本", [
        L("6401", "CNY", dr=D("1760000")),
        L("1405", "CNY", cr=D("1760000")),
    ])
    reval(3, date(2026, 3, 31))

    # ---- 4 月 ----
    V(date(2026, 4, 3), "信用证保证金退回", [
        L("1002.02", "USD", dr=D("60000")),
        L("1012", "USD", cr=D("60000")),
    ])
    V(date(2026, 4, 8), "出口销售-美国 ACME 公司", [
        L("1122.02", "USD", dr=D("165000")),
        L("6001.01", "USD", cr=D("165000")),
    ])
    # [预埋点1] 本位币按 6.80 折算(正确应按 7.12, 差异 -28,800)
    VA(date(2026, 4, 11), "收到 ACME 公司货款", [
        L("1002.02", "USD", dr=D("90000"), rate=ANOM_RATE),
        L("1122.02", "USD", cr=D("90000"), rate=ANOM_RATE),
    ], {"1002.02", "1122.02"})
    V(date(2026, 4, 15), "结汇 150,000 美元(结汇价 7.11)", [
        L("1002.01", "CNY", dr=D("1066500")),
        L("6603.02", "CNY", dr=D("1500")),
        L("1002.02", "USD", cr=D("150000")),
    ])
    V(date(2026, 4, 18), "采购钢材(13% 增值税)", [
        L("1403.01", "CNY", dr=D("380000")),
        L("2221.02", "CNY", dr=D("49400")),
        L("2202.01", "CNY", cr=D("429400")),
    ])
    V(date(2026, 4, 22), "出口销售-香港中贸公司", [
        L("1122.04", "HKD", dr=D("420000")),
        L("6001.01", "HKD", cr=D("420000")),
    ])
    V(date(2026, 4, 25), "收到香港中贸公司货款", [
        L("1002.04", "HKD", dr=D("260000")),
        L("1122.04", "HKD", cr=D("260000")),
    ])
    V(date(2026, 4, 28), "计提 4 月工资", [
        L("6601.01", "CNY", dr=D("68000")),
        L("6602.01", "CNY", dr=D("102000")),
        L("2211", "CNY", cr=D("170000")),
    ])
    V(date(2026, 4, 28), "发放 4 月工资", [
        L("2211", "CNY", dr=D("170000")),
        L("1002.01", "CNY", cr=D("170000")),
    ])
    V(date(2026, 4, 30), "计提折旧", [
        L("6602.02", "CNY", dr=D("28500")),
        L("1602", "CNY", cr=D("28500")),
    ])
    V(date(2026, 4, 30), "支付出口海运费", [
        L("6601.02", "CNY", dr=D("24000")),
        L("1002.01", "CNY", cr=D("24000")),
    ])
    V(date(2026, 4, 30), "采购产成品(13% 增值税)", [
        L("1405", "CNY", dr=D("1380000")),
        L("2221.02", "CNY", dr=D("179400")),
        L("2202.01", "CNY", cr=D("1559400")),
    ])
    V(date(2026, 4, 30), "结转 4 月销售成本", [
        L("6401", "CNY", dr=D("1420000")),
        L("1405", "CNY", cr=D("1420000")),
    ])
    reval(4, date(2026, 4, 30))

    # ---- 5 月 ----
    V(date(2026, 5, 6), "进口标准件(赊购, 美元结算)", [
        L("1403.02", "CNY", dr=D("854400")),
        L("2202.02", "USD", cr=D("120000")),
    ])
    V(date(2026, 5, 7), "出口销售-美国 ACME 公司", [
        L("1122.02", "USD", dr=D("195000")),
        L("6001.01", "USD", cr=D("195000")),
    ])
    # [预埋点1] 本位币按 6.80 折算(正确应按 7.12, 差异 -19,200)
    VA(date(2026, 5, 12), "支付美元供应商 Norton 货款", [
        L("2202.02", "USD", dr=D("60000"), rate=ANOM_RATE),
        L("1002.02", "USD", cr=D("60000"), rate=ANOM_RATE),
    ], {"2202.02", "1002.02"})
    V(date(2026, 5, 15), "收到德国 Kraft 公司货款", [
        L("1002.03", "EUR", dr=D("80000")),
        L("1122.03", "EUR", cr=D("80000")),
    ])
    V(date(2026, 5, 19), "结汇 180,000 美元(结汇价 7.14)", [
        L("1002.01", "CNY", dr=D("1285200")),
        L("1002.02", "USD", cr=D("180000")),
        L("6603.02", "CNY", cr=D("3600")),
    ])
    V(date(2026, 5, 21), "出口销售-法国 SNR 公司", [
        L("1122.03", "EUR", dr=D("95000")),
        L("6001.01", "EUR", cr=D("95000")),
    ])
    V(date(2026, 5, 26), "支付国内供应商货款", [
        L("2202.01", "CNY", dr=D("400000")),
        L("1002.01", "CNY", cr=D("400000")),
    ])
    V(date(2026, 5, 28), "计提 5 月工资", [
        L("6601.01", "CNY", dr=D("68000")),
        L("6602.01", "CNY", dr=D("102000")),
        L("2211", "CNY", cr=D("170000")),
    ])
    V(date(2026, 5, 28), "发放 5 月工资", [
        L("2211", "CNY", dr=D("170000")),
        L("1002.01", "CNY", cr=D("170000")),
    ])
    V(date(2026, 5, 29), "报销办公费(现金)", [
        L("6602.03", "CNY", dr=D("7200")),
        L("1001", "CNY", cr=D("7200")),
    ])
    V(date(2026, 5, 31), "计提折旧", [
        L("6602.02", "CNY", dr=D("28500")),
        L("1602", "CNY", cr=D("28500")),
    ])
    V(date(2026, 5, 31), "采购产成品(13% 增值税)", [
        L("1405", "CNY", dr=D("1720000")),
        L("2221.02", "CNY", dr=D("223600")),
        L("2202.01", "CNY", cr=D("1943600")),
    ])
    V(date(2026, 5, 31), "结转 5 月销售成本", [
        L("6401", "CNY", dr=D("1610000")),
        L("1405", "CNY", cr=D("1610000")),
    ])
    reval(5, date(2026, 5, 31))

    # ---- 6 月 ----
    V(date(2026, 6, 4), "购汇 100,000 欧元(购汇价 7.94)", [
        L("1002.03", "EUR", dr=D("100000")),
        L("6603.02", "CNY", dr=D("11000")),
        L("1002.01", "CNY", cr=D("794000")),
    ])
    V(date(2026, 6, 8), "出口销售-美国 ACME 公司", [
        L("1122.02", "USD", dr=D("240000")),
        L("6001.01", "USD", cr=D("240000")),
    ])
    V(date(2026, 6, 12), "收到 ACME 公司货款", [
        L("1002.02", "USD", dr=D("200000")),
        L("1122.02", "USD", cr=D("200000")),
    ])
    V(date(2026, 6, 16), "支付美元供应商 Norton 货款", [
        L("2202.02", "USD", dr=D("90000")),
        L("1002.02", "USD", cr=D("90000")),
    ])
    V(date(2026, 6, 18), "出口销售-香港中贸公司", [
        L("1122.04", "HKD", dr=D("380000")),
        L("6001.01", "HKD", cr=D("380000")),
    ])
    V(date(2026, 6, 19), "在建工程转固", [
        L("1601", "CNY", dr=D("360000")),
        L("1604", "CNY", cr=D("360000")),
    ])
    V(date(2026, 6, 22), "收到香港中贸公司货款", [
        L("1002.04", "HKD", dr=D("320000")),
        L("1122.04", "HKD", cr=D("320000")),
    ])
    V(date(2026, 6, 24), "支付出口海运费", [
        L("6601.02", "CNY", dr=D("28000")),
        L("1002.01", "CNY", cr=D("28000")),
    ])
    V(date(2026, 6, 25), "境内销售机械配件(13% 增值税)", [
        L("1122.01", "CNY", dr=D("293800")),
        L("6001.02", "CNY", cr=D("260000")),
        L("2221.01", "CNY", cr=D("33800")),
    ])
    V(date(2026, 6, 26), "出售废料(13% 增值税)", [
        L("1002.01", "CNY", dr=D("20340")),
        L("6051", "CNY", cr=D("18000")),
        L("2221.01", "CNY", cr=D("2340")),
    ])
    V(date(2026, 6, 30), "支付二季度贷款利息", [
        L("6603.01", "CNY", dr=D("42000")),
        L("1002.01", "CNY", cr=D("42000")),
    ])
    V(date(2026, 6, 30), "计提 6 月工资", [
        L("6601.01", "CNY", dr=D("68000")),
        L("6602.01", "CNY", dr=D("102000")),
        L("2211", "CNY", cr=D("170000")),
    ])
    V(date(2026, 6, 30), "发放 6 月工资", [
        L("2211", "CNY", dr=D("170000")),
        L("1002.01", "CNY", cr=D("170000")),
    ])
    V(date(2026, 6, 30), "计提折旧", [
        L("6602.02", "CNY", dr=D("28500")),
        L("1602", "CNY", cr=D("28500")),
    ])
    V(date(2026, 6, 30), "采购产成品(13% 增值税)", [
        L("1405", "CNY", dr=D("2030000")),
        L("2221.02", "CNY", dr=D("263900")),
        L("2202.01", "CNY", cr=D("2293900")),
    ])
    V(date(2026, 6, 30), "结转 6 月销售成本", [
        L("6401", "CNY", dr=D("1950000")),
        L("1405", "CNY", cr=D("1950000")),
    ])
    reval(6, date(2026, 6, 30))


build()

# ---------------------------------------------------------------- TB 组装
# 预埋点2: TB 期初覆写(正确口径 -> 虚增口径)
TB_OPEN_OVERRIDE = {
    ("1002.02", "USD"): (OPEN_SNAP_O[("1002.02", "USD")] + OPEN_DIFF_USD,
                         OPEN_SNAP_B[("1002.02", "USD")] + OPEN_DIFF_CNY),
    ("4104", "CNY"): (PLUG_4104 - OPEN_DIFF_CNY, PLUG_4104 - OPEN_DIFF_CNY),
}


def build_tb_rows():
    agg = defaultdict(lambda: [ZERO, ZERO, ZERO, ZERO])  # 原币借/贷, 本位币借/贷
    for e in entries:
        k = (e["subj"], e["cur"])
        a = agg[k]
        a[0] += e["od"] or ZERO
        a[1] += e["oc"] or ZERO
        a[2] += e["bd"] or ZERO
        a[3] += e["bc"] or ZERO
    rows = []
    for k in sorted(set(OPEN_SNAP_O) | set(agg), key=lambda x: (x[0], CUR_ORDER[x[1]])):
        s, c = k
        o0, b0 = OPEN_SNAP_O.get(k, ZERO), OPEN_SNAP_B.get(k, ZERO)
        fo = agg.get(k)
        s1o = bal_orig.get(k, ZERO)
        s1b = bal_base.get(k, ZERO)
        if not fo and o0 == 0 and b0 == 0:
            continue
        wo, wb = TB_OPEN_OVERRIDE.get(k, (o0, b0))       # 写入 TB 的期初
        rows.append({
            "subj": s, "name": SUBJ_NAME[s], "cur": CUR_CN[c], "curcode": c,
            "open": (wo, wb), "open_correct": (o0, b0),
            "flow": (fo or [ZERO, ZERO, ZERO, ZERO]),
            "end": (s1o, s1b),
        })
    return rows


TB_ROWS = build_tb_rows()

# ---------------------------------------------------------------- 写 Excel
F_TITLE = Font(name="宋体", size=14, bold=True)
F_INFO = Font(name="宋体", size=9)
F_HEAD = Font(name="宋体", size=10, bold=True)
F_BODY = Font(name="宋体", size=10)
A_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_LEFT = Alignment(horizontal="left", vertical="center")
AMT_FMT = "#,##0.00"
DATE_FMT = "yyyy-mm-dd"


def put_amt(ws, r, c, v):
    if v is None or v == 0:
        return
    cell = ws.cell(row=r, column=c)
    cell.value = float(v)
    cell.number_format = AMT_FMT
    cell.font = F_BODY


def write_tb(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "科目余额表"
    ws["A1"] = "科目余额表"
    ws["A1"].font = F_TITLE
    ws["A1"].alignment = A_CENTER
    ws.merge_cells("A1:O1")
    ws["A2"] = ("单位：%s    期间：2026年第1-6期(2026-01-01至2026-06-30)    "
                "币别：所有币别(综合本位币)    制表：金蝶K/3 Cloud" % ENTITY)
    ws["A2"].font = F_INFO
    ws["A2"].alignment = A_CENTER
    ws.merge_cells("A2:O2")
    # 第 3-4 行复合表头
    for col, text in ((1, "科目编码"), (2, "科目名称"), (3, "币种")):
        ws.cell(row=3, column=col, value=text)
        ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
    groups = [(4, "期初余额"), (8, "本期发生"), (12, "期末余额")]
    sub = ["借方(原币金额)", "借方(本位币金额)", "贷方(原币金额)", "贷方(本位币金额)"]
    for start, title in groups:
        ws.cell(row=3, column=start, value=title)
        ws.merge_cells(start_row=3, start_column=start, end_row=3, end_column=start + 3)
        for i, label in enumerate(sub):
            cell = ws.cell(row=4, column=start + i, value=label)
            cell.font = F_HEAD
            cell.alignment = A_CENTER
    for col in (1, 2, 3):
        cell = ws.cell(row=3, column=col)
        cell.font = F_HEAD
        cell.alignment = A_CENTER
    ws.cell(row=3, column=4).font = F_HEAD
    ws.cell(row=3, column=4).alignment = A_CENTER
    ws.cell(row=3, column=8).font = F_HEAD
    ws.cell(row=3, column=8).alignment = A_CENTER
    ws.cell(row=3, column=12).font = F_HEAD
    ws.cell(row=3, column=12).alignment = A_CENTER

    r = 5
    for row in TB_ROWS:
        wo, wb_ = row["open"]
        fo_d, fo_c, fb_d, fb_c = row["flow"]
        eo, eb = row["end"]
        ws.cell(row=r, column=1, value=row["subj"]).font = F_BODY
        ws.cell(row=r, column=2, value=row["name"]).font = F_BODY
        ws.cell(row=r, column=3, value=row["cur"]).font = F_BODY
        # 期初(按写入方向拆借/贷)
        if wo >= 0:
            put_amt(ws, r, 4, wo); put_amt(ws, r, 5, wb_)
        else:
            put_amt(ws, r, 6, -wo); put_amt(ws, r, 7, -wb_)
        put_amt(ws, r, 8, fo_d); put_amt(ws, r, 9, fb_d)
        put_amt(ws, r, 10, fo_c); put_amt(ws, r, 11, fb_c)
        if eo >= 0:
            put_amt(ws, r, 12, eo); put_amt(ws, r, 13, eb)
        else:
            put_amt(ws, r, 14, -eo); put_amt(ws, r, 15, -eb)
        r += 1

    widths = {1: 10, 2: 30, 3: 9, 4: 15, 5: 15, 6: 15, 7: 15, 8: 15, 9: 15,
              10: 15, 11: 15, 12: 15, 13: 15, 14: 15, 15: 15}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "D5"
    wb.save(path)


def write_je(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "序时账"
    ws["A1"] = "序时账(核算项目明细账)"
    ws["A1"].font = F_TITLE
    ws["A1"].alignment = A_CENTER
    ws.merge_cells("A1:M1")
    ws["A2"] = ("单位：%s    期间：2026年第1-6期(2026-01-01至2026-06-30)    "
                "币别：所有币别(综合本位币)    制表：金蝶K/3 Cloud" % ENTITY)
    ws["A2"].font = F_INFO
    ws["A2"].alignment = A_CENTER
    ws.merge_cells("A2:M2")
    headers = ["日期", "凭证字", "凭证号", "摘要", "科目编码", "科目名称", "币种",
               "借方(原币)", "贷方(原币)", "借方(本位币)", "贷方(本位币)",
               "余额(原币)", "余额(本位币)"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=i, value=h)
        cell.font = F_HEAD
        cell.alignment = A_CENTER
    # 余额列: 按科目 x 币种自(正确)期初起算的连续余额, 借正贷负
    run_o = dict(OPEN_SNAP_O)
    run_b = dict(OPEN_SNAP_B)
    r = 4
    for e in entries:
        k = (e["subj"], e["cur"])
        run_o[k] = run_o.get(k, ZERO) + (e["od"] or ZERO) - (e["oc"] or ZERO)
        run_b[k] = run_b.get(k, ZERO) + (e["bd"] or ZERO) - (e["bc"] or ZERO)
        cell = ws.cell(row=r, column=1, value=e["date"])
        cell.number_format = DATE_FMT
        cell.font = F_BODY
        ws.cell(row=r, column=2, value=e["word"]).font = F_BODY
        ws.cell(row=r, column=3, value=e["no"]).font = F_BODY
        ws.cell(row=r, column=4, value=e["summ"]).font = F_BODY
        ws.cell(row=r, column=5, value=e["subj"]).font = F_BODY
        ws.cell(row=r, column=6, value=SUBJ_NAME[e["subj"]]).font = F_BODY
        ws.cell(row=r, column=7, value=CUR_CN[e["cur"]]).font = F_BODY
        put_amt(ws, r, 8, e["od"])
        put_amt(ws, r, 9, e["oc"])
        put_amt(ws, r, 10, e["bd"])
        put_amt(ws, r, 11, e["bc"])
        put_amt(ws, r, 12, run_o[k])
        put_amt(ws, r, 13, run_b[k])
        r += 1
    widths = {1: 11, 2: 7, 3: 7, 4: 34, 5: 10, 6: 30, 7: 9, 8: 14, 9: 14,
              10: 14, 11: 14, 12: 14, 13: 14}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "E4"
    wb.save(path)


write_tb(TB_PATH)
write_je(JE_PATH)

# ---------------------------------------------------------------- 自检
CHECKS = []


def check(name, ok, detail=""):
    CHECKS.append((name, bool(ok), detail))
    tag = "[通过] " if ok else "[失败] "
    line = tag + name
    if detail:
        line += "  -- " + detail
    print(line)


def dec(v):
    if v is None or v == "":
        return ZERO
    return Decimal(str(v)).quantize(TWO, rounding=ROUND_HALF_UP)


def load_tb():
    ws = load_workbook(TB_PATH).active
    head3 = [ws.cell(row=3, column=c).value for c in range(1, 16)]
    head4 = [ws.cell(row=4, column=c).value for c in range(1, 16)]
    rows = []
    r = 5
    while ws.cell(row=r, column=1).value:
        rows.append({
            "subj": ws.cell(row=r, column=1).value,
            "name": ws.cell(row=r, column=2).value,
            "cur": ws.cell(row=r, column=3).value,
            "o_dr_o": dec(ws.cell(row=r, column=4).value), "o_dr_b": dec(ws.cell(row=r, column=5).value),
            "o_cr_o": dec(ws.cell(row=r, column=6).value), "o_cr_b": dec(ws.cell(row=r, column=7).value),
            "f_dr_o": dec(ws.cell(row=r, column=8).value), "f_dr_b": dec(ws.cell(row=r, column=9).value),
            "f_cr_o": dec(ws.cell(row=r, column=10).value), "f_cr_b": dec(ws.cell(row=r, column=11).value),
            "e_dr_o": dec(ws.cell(row=r, column=12).value), "e_dr_b": dec(ws.cell(row=r, column=13).value),
            "e_cr_o": dec(ws.cell(row=r, column=14).value), "e_cr_b": dec(ws.cell(row=r, column=15).value),
        })
        r += 1
    return head3, head4, rows


def load_je():
    ws = load_workbook(JE_PATH).active
    head = [ws.cell(row=3, column=c).value for c in range(1, 14)]
    idx = {h: i + 1 for i, h in enumerate(head)}
    rows = []
    r = 4
    while ws.cell(row=r, column=1).value is not None:
        row = {h: ws.cell(row=r, column=idx[h]).value for h in head}
        row["_row"] = r
        rows.append(row)
        r += 1
    return head, idx, rows


def self_check():
    from datetime import datetime as _dt
    print("=" * 72)
    print("金蝶(汇率损益测试集) 数据自检")
    print("=" * 72)

    # 1. 文件生成
    check("文件生成", TB_PATH.exists() and JE_PATH.exists(),
          "金蝶_科目余额表.xlsx / 金蝶_序时账.xlsx")

    head3, head4, tb = load_tb()
    head, idx, je = load_je()

    # 2. TB 表头结构: 标题/信息行/两行复合表头, 币种列与原币/本位币列各 6 组
    ok = (head3[0] == "科目编码" and head3[2] == "币种"
          and head3[3] == "期初余额" and head3[7] == "本期发生"
          and head3[11] == "期末余额")
    n_orig = sum(1 for h in head4 if h and "原币金额" in h)
    n_base = sum(1 for h in head4 if h and "本位币金额" in h)
    check("TB 复合表头(期初/本期发生/期末 x 借贷 x 原币/本位币)",
          ok and n_orig == 6 and n_base == 6,
          "原币列 %d 个, 本位币列 %d 个" % (n_orig, n_base))

    # 3. 列识别与数据形态: 币种列逐行非空, 原币/本位币列非全空
    cur_col = idx["币种"]
    cur_filled = sum(1 for r in je if r["币种"])
    orig_cnt = sum(1 for r in je if (r["借方(原币)"] is not None and r["借方(原币)"] != "")
                   or (r["贷方(原币)"] is not None and r["贷方(原币)"] != ""))
    base_cnt = sum(1 for r in je if (r["借方(本位币)"] is not None and r["借方(本位币)"] != "")
                   or (r["贷方(本位币)"] is not None and r["贷方(本位币)"] != ""))
    check("币种列/原币列/本位币列存在且非全空",
          cur_filled == len(je) and orig_cnt > 100 and base_cnt > 150,
          "JE %d 行全部有币种, 原币有值 %d 行, 本位币有值 %d 行" % (len(je), orig_cnt, base_cnt))

    # 4. 每凭证本位币借贷平衡
    vbal = defaultdict(lambda: [ZERO, ZERO])
    for r in je:
        key = (r["凭证字"], r["凭证号"])
        vbal[key][0] += dec(r["借方(本位币)"])
        vbal[key][1] += dec(r["贷方(本位币)"])
    bad = [k for k, (a, b) in vbal.items() if a != b]
    check("每凭证本位币借贷平衡", not bad,
          "共 %d 张凭证, 不平衡 %d 张" % (len(vbal), len(bad)))

    # 5. 凭证号连续唯一
    nos = [r["凭证号"] for r in je]
    uniq = sorted(set(nos))
    check("凭证号唯一且连续", len(uniq) == len(set(nos)) and uniq == list(range(1, len(uniq) + 1)),
          "凭证号 1..%d" % max(nos))

    # 6. TB 滚动勾稽(除预埋点2 涉及的两个科目)
    je_net = defaultdict(lambda: [ZERO, ZERO])
    for r in je:
        key = (r["科目编码"], r["币种"])
        je_net[key][0] += dec(r["借方(原币)"]) - dec(r["贷方(原币)"])
        je_net[key][1] += dec(r["借方(本位币)"]) - dec(r["贷方(本位币)"])
    diffs = {}
    for t in tb:
        key = (t["subj"], t["cur"])
        s0o = t["o_dr_o"] - t["o_cr_o"]
        s0b = t["o_dr_b"] - t["o_cr_b"]
        s1o = t["e_dr_o"] - t["e_cr_o"]
        s1b = t["e_dr_b"] - t["e_cr_b"]
        do = s1o - s0o - je_net.get(key, [ZERO, ZERO])[0]
        db = s1b - s0b - je_net.get(key, [ZERO, ZERO])[1]
        if do != 0 or db != 0:
            diffs[key] = (do, db)
    expect = {("1002.02", "美元"): (-OPEN_DIFF_USD, -OPEN_DIFF_CNY),
              # 4104 为人民币行, 原币列与本位币列同值镜像, 两列差异均为 +356,000
              ("4104", "人民币"): (OPEN_DIFF_CNY, OPEN_DIFF_CNY)}
    ok = diffs == expect
    detail = "勾稽差异科目 %d 个(预期 2 个, 均为预埋点2)" % len(diffs)
    check("TB 滚动勾稽: 期末=期初+JE净发生(除预埋点)", ok, detail)
    for k, (do, db) in sorted(diffs.items()):
        print("       预期差异 %s %s: 原币 %s, 本位币 %s" % (k[0], k[1], fmt(do), fmt(db)))

    # 7. TB 总额借贷平衡(期初/本期发生/期末, 本位币)
    t0 = sum(t["o_dr_b"] for t in tb) - sum(t["o_cr_b"] for t in tb)
    tf = sum(t["f_dr_b"] for t in tb) - sum(t["f_cr_b"] for t in tb)
    t1 = sum(t["e_dr_b"] for t in tb) - sum(t["e_cr_b"] for t in tb)
    check("TB 期初/发生/期末总额借贷平衡(本位币)",
          t0 == 0 and tf == 0 and t1 == 0,
          "期初差 %s, 发生差 %s, 期末差 %s" % (fmt(t0), fmt(tf), fmt(t1)))

    # 8. 规模
    check("数据规模(TB 40-80 行, JE 150-300 行)",
          40 <= len(tb) <= 80 and 150 <= len(je) <= 300,
          "TB %d 行, JE %d 行, 凭证 %d 张" % (len(tb), len(je), len(vbal)))

    # 9. 金额/日期单元格类型
    amt_ok = True
    for r in je:
        if not isinstance(r["日期"], (_dt, date)):
            amt_ok = False
        for h in ("借方(原币)", "贷方(原币)", "借方(本位币)", "贷方(本位币)", "余额(原币)", "余额(本位币)"):
            v = r[h]
            if v is not None and not isinstance(v, (int, float)):
                amt_ok = False
    check("JE 金额为数值型且日期为日期型", amt_ok)

    # 10. 预埋点1: 美元分录隐含汇率扫描
    hits = []
    for r in je:
        if r["币种"] != "美元":
            continue
        for d_h, b_h in (("借方(原币)", "借方(本位币)"), ("贷方(原币)", "贷方(本位币)")):
            o, b = dec(r[d_h]), dec(r[b_h])
            if o > 0 and b > 0:
                rate = (b / o).quantize(Decimal("0.0001"))
                if abs(rate - BOOK["USD"]) > Decimal("0.005"):
                    hits.append((r["凭证号"], r["日期"], r["科目编码"], o, b, rate))
    ok = len(hits) == 6 and all(h[5] == D("6.8000") for h in hits)
    check("预埋点1: 隐含汇率异常行恰为 3 张凭证/6 行(全部 6.80)", ok,
          "扫描到 %d 行" % len(hits))
    for h in hits:
        print("       记-%d %s %s 原币 %s 本位币 %s 隐含汇率 %s"
              % (h[0], h[1].strftime("%Y-%m-%d"), h[2], fmt(h[3]), fmt(h[4]), h[5]))

    # 11. 干净基准: 外币货币性科目期末余额 = 期末原币 x 期末汇率
    bad_end = []
    for t in tb:
        code = t["subj"]
        if code not in MONETARY_FX:
            continue
        s1o = t["e_dr_o"] - t["e_cr_o"]
        s1b = t["e_dr_b"] - t["e_cr_b"]
        expect_b = r2(s1o * END_RATE[{"美元": "USD", "欧元": "EUR", "港元": "HKD", "日元": "JPY"}[t["cur"]]])
        if s1b != expect_b:
            bad_end.append((code, t["cur"], s1o, s1b, expect_b))
    check("外币货币性项目期末余额=期末原币x期末汇率(重估完整)", not bad_end,
          "检查 %d 个科目x币种" % sum(1 for t in tb if t["subj"] in MONETARY_FX))

    # 汇总
    failed = [c for c in CHECKS if not c[1]]
    print("-" * 72)
    print("自检汇总: %d 项检查, 通过 %d 项, 失败 %d 项" % (len(CHECKS), len(CHECKS) - len(failed), len(failed)))
    print("TB %d 行 | JE %d 行 | 凭证 %d 张" % (len(tb), len(je), len(vbal)))
    dr6603 = sum(dec(r["借方(本位币)"]) for r in je if r["科目编码"] == "6603.02")
    cr6603 = sum(dec(r["贷方(本位币)"]) for r in je if r["科目编码"] == "6603.02")
    print("6603.02 汇兑损益: 借方(损失) %s, 贷方(收益) %s, 净额 %s"
          % (fmt(dr6603), fmt(cr6603), fmt(cr6603 - dr6603)))
    return len(failed) == 0


def report_points():
    print("-" * 72)
    print("预埋测试点核对(供 README 引用)")
    for no, d, subj, summ, orig, base, rate in anomaly_rows:
        print("  [点1] 记-%d %s %s (%s): 原币 %s, 本位币 %s, 隐含汇率 %s (记账汇率 7.12)"
              % (no, d.strftime("%Y-%m-%d"), subj, summ, fmt(orig), fmt(base), rate))
    row = next(t for t in TB_ROWS if t["subj"] == "1002.02" and t["curcode"] == "USD")
    co, cb = row["open_correct"]
    wo, wb_ = row["open"]
    print("  [点2] 1002.02 美元户期初: 表列 %s USD / %s CNY, 正确值 %s USD / %s CNY, 虚增 %s / %s"
          % (fmt(wo), fmt(wb_), fmt(co), fmt(cb), fmt(wo - co), fmt(wb_ - cb)))
    print("  [点2] 4104 未分配利润期初(贷方): 表列 %s, 正确值 %s, 虚增 %s"
          % (fmt(-TB_OPEN_OVERRIDE[("4104", "CNY")][0]), fmt(-PLUG_4104), fmt(OPEN_DIFF_CNY)))


if __name__ == "__main__":
    ok = self_check()
    report_points()
    sys.exit(0 if ok else 1)
